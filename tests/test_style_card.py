"""T-029: style learned from example articles, truth kept by a fact sheet and two separate reviews (offline)."""
import json
import tempfile
import unittest
from pathlib import Path

import test_core
from test_article_quality import SequenceTransport, answers, chat, system_text, user_json
from wewriter.api import WRITER_RULES, ApiError, Model
from wewriter.config import defaults, migrate
from wewriter.db import SCHEMA_VERSION, Database
from wewriter.pipeline.composition import style_pass, use_style
from wewriter.style import (CARD_MAX, DEFAULT_STYLE_CARD, SKELETON_MAX, card_text, choose_example, clean_style_card,
                            fact_sheet,
                            cite_facts, read_example_file, same_facts, skeleton, strength_problems, style_block,
                            valid_style_card)
from wewriter.webapp import Api, Backend

FIXTURE = Path(__file__).parent / 'fixtures' / 'style' / 'example_fictional.md'
SOURCES = [
    {'source_id': 'S1', 'article_id': 1, 'title': '验收要求', 'account_name': '示例号', 'url': 'https://mp.weixin.qq.com/s/a',
     'analysis': [{'summary': '示例市验收要求', 'facts': [
         {'claim': '新设备须连续稳定试运行不少于30天', 'quote': '连续稳定试运行不少于30天'},
         {'claim': '严禁大额现金支付设备款', 'quote': '严禁出现以拨代支、大额现金支付设备款'}],
         'uncertainties': ['材料未注明发文机关']}],
     'web_research': '国家文件取消了1亿元门槛[W1]。示例市公告列出验收名单[W2]。未找到官方出处：2090万元。'},
    {'source_id': 'W1', 'title': '国家通知', 'url': 'https://www.ndrc.gov.cn/notice', 'kind': 'web'},
    {'source_id': 'W2', 'title': '行业网站转载', 'url': 'https://news.example.com/a', 'kind': 'web'},
]
CONFIG = {'topic': '', 'style': '', 'target_words': 1500, 'writing_mode': 'single_web'}
ARTICLE = {'title': '示例市设备验收：跑满30天再报', 'outline': [], 'review_notes': [],
           'body': '示例市的设备验收开始了，新设备要连续试运行30天。[S1]\n\n国家已取消1亿元门槛。[W1]'}


class FactSheetTest(unittest.TestCase):
    def test_web_marks_after_punctuation_stay_with_their_facts(self):
        sources = [
            {'source_id': 'S1', 'web_research': '第一项现行规定。[W1] 第二项另有要求。 [W2]\n没有出处的句子。'},
            {'source_id': 'W1', 'title': '示例通知甲', 'url': 'https://www.ndrc.gov.cn/example-a'},
            {'source_id': 'W2', 'title': '示例通知乙', 'url': 'https://www.ndrc.gov.cn/example-b'},
        ]
        facts = fact_sheet(sources)
        self.assertEqual([(f['fact'], f['cite']) for f in facts],
                         [('第一项现行规定。', '[W1]'), ('第二项另有要求。', '[W2]')])

    def test_levels_scope_and_strength(self):
        facts = fact_sheet(SOURCES)
        by_text = {f['fact']: f for f in facts}
        trial = by_text['新设备须连续稳定试运行不少于30天']
        self.assertEqual((trial['level'], trial['sources']), ('reported', ['S1']))
        self.assertIn('行业资料', trial['scope'])
        self.assertNotIn('示例号', json.dumps(facts, ensure_ascii=False))     # iron rule: no account names
        self.assertIn('不能写成“官方明确”', trial['strength'])
        self.assertEqual(by_text['材料未注明发文机关']['level'], 'unverified')
        self.assertEqual(by_text['国家文件取消了1亿元门槛。']['level'], 'official')        # ndrc.gov.cn
        self.assertEqual(by_text['示例市公告列出验收名单。']['level'], 'reported')        # a non-government site
        self.assertEqual(by_text['未找到官方出处：2090万元。']['level'], 'unverified')
        self.assertEqual([f['id'] for f in facts], [f'F{n}' for n in range(1, len(facts) + 1)])

    def test_wording_stronger_than_the_source(self):
        self.assertEqual(strength_problems('新设备必须试运行30天，否则直接驳回。[S1]', SOURCES),
                         ['说法强度超过来源可信等级（必须、直接驳回）：新设备必须试运行30天，否则直接驳回。[S1]'])
        self.assertEqual(strength_problems('严禁大额现金支付设备款。[S1]', SOURCES), [])          # the source says 严禁
        self.assertEqual(strength_problems('项目必须按国家通知办理。[W1]', SOURCES), [])          # official source
        self.assertTrue(strength_problems('官方明确要求一律提交10类资料。[W2]', SOURCES))


class StyleInputTest(unittest.TestCase):
    def test_customer_approved_fallback_is_only_a_writing_method(self):
        transport = SequenceTransport(chat(ARTICLE))
        model = Model('deepseek', defaults()['providers']['deepseek'], 'test-only', transport)
        model.write(SOURCES, dict(CONFIG, style='像同行解释，标题具体但不能夸大'))
        sent = system_text(transport.payloads[0])
        user = user_json(transport.payloads[0])
        self.assertIn('从 facts 选一个读者需要作出的判断或动作', sent)
        self.assertIn('正文篇幅应达到 target_chinese_characters 的 90% 至 110%', sent)
        self.assertIn('【风格卡】', sent)
        self.assertIn('申报这类项目，书面申请要交上去', sent)
        self.assertNotIn('遵义', sent)
        self.assertNotIn('马鞍山', sent)
        self.assertEqual(user['user_style'], '像同行解释，标题具体但不能夸大')
        self.assertTrue(all(f['fact'] not in sent for f in user['facts']))

    def test_configured_card_replaces_fallback_without_becoming_evidence(self):
        card = dict(DEFAULT_STYLE_CARD, opening='先用一个经证实的限制说明读者为何要读')
        transport = SequenceTransport(chat(ARTICLE))
        model = Model('deepseek', defaults()['providers']['deepseek'], 'test-only', transport)
        model.write(SOURCES, dict(CONFIG, style_card=card, style='开头直接讲事实'))
        sent = system_text(transport.payloads[0])
        self.assertIn(card['opening'], sent)
        self.assertNotIn(DEFAULT_STYLE_CARD['opening'], sent)
        self.assertIn('user_style 高于已配置风格卡', sent)
        self.assertEqual(user_json(transport.payloads[0])['user_style'], '开头直接讲事实')

    def test_card_and_skeleton_sizes(self):
        self.assertTrue(valid_style_card(DEFAULT_STYLE_CARD))
        self.assertLessEqual(len(card_text(DEFAULT_STYLE_CARD)), CARD_MAX)
        self.assertGreaterEqual(len(card_text(DEFAULT_STYLE_CARD)), 400)
        title, body, meta = read_example_file(FIXTURE)
        self.assertEqual((title[:8], meta['score'], meta['kind']), ('示例补贴验收开始', '80', 'good'))
        shape = skeleton(title, body * 5, meta['avoid'])
        self.assertLessEqual(len(shape), SKELETON_MAX)
        self.assertIn('开头：最近跑项目的朋友要盯紧了', shape)
        self.assertIn('不要学：', shape)

    def test_budget_cuts_the_example_and_says_so(self):
        example = {'skeleton': '甲' * 900}
        text, note = style_block(DEFAULT_STYLE_CARD, example, budget=1200)
        self.assertLessEqual(len(text), 1200)
        self.assertIn('已截短范文骨架', note)
        text, note = style_block(DEFAULT_STYLE_CARD, example)
        self.assertEqual(note, '')

    def test_prefix_is_identical_across_articles_and_the_library_size_does_not_grow_the_input(self):
        db = Database(Path(tempfile.mkdtemp()) / 'x.sqlite')
        title, body, meta = read_example_file(FIXTURE)
        db.add_example(title, body, 'good', 80, meta['avoid'])

        def prompts():
            out = []
            for article in ('示例补贴验收', '另一篇完全不同的稿件'):
                config = dict(CONFIG)
                use_style(db, config, article)
                transport = SequenceTransport(chat(ARTICLE))
                Model('deepseek', defaults()['providers']['deepseek'], 'k', transport).write(
                    [dict(SOURCES[0], title=article)] + SOURCES[1:], config)
                out.append(system_text(transport.payloads[0]))
            return out
        first, second = prompts()
        self.assertEqual(first, second)                      # same rules, card and example: same prefix
        self.assertTrue(first.startswith(WRITER_RULES + '\n\n【风格卡】'))
        for n in range(12):
            db.add_example(f'其他范文{n}', body + '\n\n更多内容。' * 40, 'good', 60)
        grown = prompts()[0]
        self.assertLessEqual(len(grown), len(WRITER_RULES) + 2 + 2500 + 400)   # bounded by the style budget
        self.assertLessEqual(abs(len(grown) - len(first)), SKELETON_MAX)      # at most a different example

    def test_only_the_writer_carries_style(self):
        config = dict(CONFIG, style_card=DEFAULT_STYLE_CARD, example={'id': 1, 'title': 't', 'skeleton': '骨架内容'})
        revision = {**ARTICLE, 'quality_notes': [], 'claim_checks': [
            {'sentence': '标题：示例市设备验收：跑满30天再报', 'evidence': 'S1', 'action': '保留', 'note': ''}]}
        transport = SequenceTransport(chat(ARTICLE), chat(revision))
        model = Model('deepseek', defaults()['providers']['deepseek'], 'k', transport)
        model.review(model.write(SOURCES, config), SOURCES, config)
        write, review = transport.payloads
        self.assertIn('骨架内容', system_text(write))
        self.assertIn('facts', user_json(write))
        self.assertNotIn('骨架内容', system_text(review) + json.dumps(user_json(review), ensure_ascii=False))
        self.assertNotIn('【风格卡】', system_text(review))


class FactCitationTest(unittest.TestCase):
    """The article cites sources ([S1][W2]); a fact number ([F1]) is mapped back instead of failing the write."""

    def test_each_fact_carries_its_source_marks(self):
        facts = fact_sheet(SOURCES)
        self.assertTrue(all(f['cite'] == ''.join(f'[{s}]' for s in f['sources']) for f in facts))

    def test_fact_numbers_in_the_body_become_source_marks(self):
        facts = fact_sheet(SOURCES)
        trial = next(f for f in facts if f['fact'].startswith('新设备须'))
        body = f'新设备要连续试运行30天。[{trial["id"]}]\n\n国家已取消1亿元门槛。[W1]'
        transport = SequenceTransport(chat({**ARTICLE, 'body': body}))
        written = Model('deepseek', defaults()['providers']['deepseek'], 'k', transport).write(SOURCES, CONFIG)
        self.assertIn('试运行30天。[S1]', written['body'])
        self.assertNotIn('[F', written['body'])
        self.assertIn('cite', system_text(transport.payloads[0]))

    def test_an_unknown_mark_is_named_in_the_error(self):
        bad = {**ARTICLE, 'body': '新设备要连续试运行30天。[F99]\n\n国家已取消1亿元门槛。[X3]'}
        transport = SequenceTransport(chat(bad), chat(bad))
        with self.assertRaises(ApiError) as caught:
            Model('deepseek', defaults()['providers']['deepseek'], 'k', transport).write(SOURCES, CONFIG)
        self.assertIn('X3', str(caught.exception))

    def test_a_fact_number_not_on_the_sheet_never_reaches_the_article(self):
        # Seen in the 2026-10-27 comparison: “…未查到对应的中标公告原文 [S16][F12]” went out with a bare F number.
        body = '新设备要连续试运行30天。[S1][F99]\n\n国家已取消1亿元门槛。[W1][W1]'
        self.assertEqual(cite_facts(body, fact_sheet(SOURCES)), '新设备要连续试运行30天。[S1]\n\n国家已取消1亿元门槛。[W1]')
        transport = SequenceTransport(chat({**ARTICLE, 'body': body}))
        written = Model('deepseek', defaults()['providers']['deepseek'], 'k', transport).write(SOURCES, CONFIG)
        self.assertNotIn('[F', written['body'])
        revised = {**ARTICLE, 'body': ARTICLE['body'] + '\n\n示例市公告列出了验收名单。[W2][F77]', 'quality_notes': []}
        transport = SequenceTransport(chat(revised), answers(ARTICLE, revised))
        reviewed = Model('deepseek', defaults()['providers']['deepseek'], 'k', transport).review(dict(ARTICLE), SOURCES, CONFIG)
        self.assertNotIn('[F', reviewed['body'])
        self.assertNotIn('[F', json.dumps(user_json(transport.payloads[1]), ensure_ascii=False).split('"facts"')[0])


class FactReviewStrengthTest(unittest.TestCase):
    def test_a_revision_stronger_than_its_source_is_sent_back(self):
        # A changed sentence worded stronger than its source gets the targeted fix, even with a valid basis.
        strong = {**ARTICLE, 'body': ARTICLE['body'] + '\n\n跑不满30天直接驳回。[S1]', 'quality_notes': []}
        transport = SequenceTransport(chat(strong), answers(ARTICLE, strong), chat({'fixes': []}))
        result = Model('deepseek', defaults()['providers']['deepseek'], 'k', transport).review(dict(ARTICLE), SOURCES, CONFIG)
        self.assertIn('说法强度超过来源可信等级', user_json(transport.payloads[2])['sentences'][0]['problem'])
        self.assertNotIn('直接驳回', result['body'])


class StyleWriter:
    """Fake writer for style_pass: scores low and offers a rewrite."""
    rewrite = None
    calls = 0
    scores = None

    def __init__(self, *args):
        pass

    def style_review(self, article, sources, config, signals=()):
        type(self).calls += 1
        return {'scores': self.scores or {k: 2 for k in ('title', 'opening', 'hierarchy', 'voice', 'usefulness', 'distinctness')},
                'notes': ['开头太平'], 'rewrite': self.rewrite}

    def _finish_draft(self, result, sources, config):
        return Model._finish_draft(result, sources, config)


class Run:
    def __init__(self, config, writer):
        self.config, self.writer, self.styled = config, writer, None

    def stage(self, *args):
        pass


class Checker:
    @staticmethod
    def check():
        pass


class StylePassTest(unittest.TestCase):
    def setUp(self):
        StyleWriter.calls = 0
        StyleWriter.scores = None
        self.config = dict(CONFIG, recent_drafts=[{'title': ARTICLE['title'], 'opening': '', 'headings': []}])

    def run_pass(self, rewrite):
        StyleWriter.rewrite = rewrite
        run = Run(self.config, StyleWriter())
        return style_pass(Checker, run, dict(ARTICLE), SOURCES), run.styled

    def test_a_rewrite_with_the_same_facts_is_adopted(self):
        rewrite = {'title': '示例市验收开始了！新设备先跑满30天',
                   'body': '最近跑项目的朋友注意：示例市验收开始了，新设备要连续试运行30天。[S1]\n\n好消息是，国家已取消1亿元门槛。[W1]'}
        out, styled = self.run_pass(rewrite)
        self.assertEqual(out['title'], rewrite['title'])
        self.assertTrue(styled['adopted'])
        self.assertEqual(styled['scores']['opening'], 2)

    def test_a_rewrite_that_changes_facts_is_refused(self):
        for body, why in (('示例市验收开始了，新设备要连续试运行60天。[S1]\n\n国家已取消1亿元门槛。[W1]', '数字变了'),
                          ('示例市验收开始了，新设备要连续试运行30天。[S1]\n\n国家已取消1亿元门槛。[W1][W2]', '来源标记变了'),
                          ('示例市验收开始了，新设备跑不满30天直接驳回。[S1]\n\n国家已取消1亿元门槛。[W1]', '说法强度')):
            with self.subTest(why=why):
                out, styled = self.run_pass({'title': ARTICLE['title'], 'body': body})
                self.assertEqual(out['body'], ARTICLE['body'])
                self.assertFalse(styled['adopted'])
                self.assertIn(why, styled['reason'])

    def test_no_signals_still_reviews_title_opening_and_hierarchy(self):
        self.config['recent_drafts'] = []
        out, styled = self.run_pass(None)
        self.assertEqual(StyleWriter.calls, 1)
        self.assertEqual(styled['signals'], [])
        self.assertEqual(styled['scores']['opening'], 2)
        self.assertIn('待人工复核', styled['reason'])

    def test_disabled_style_review_does_not_call_model(self):
        self.config.update(recent_drafts=[], style_review_enabled=False)
        out, styled = self.run_pass(None)
        self.assertEqual(StyleWriter.calls, 0)
        self.assertIn('已关闭', styled['reason'])

    def test_core_score_three_needs_rewrite_even_without_readability_signal(self):
        self.config['recent_drafts'] = []
        StyleWriter.scores = dict.fromkeys(('title', 'opening', 'hierarchy', 'voice'), 4)
        StyleWriter.scores.update(usefulness=3, distinctness=3, opening=3)
        out, styled = self.run_pass(None)
        self.assertEqual(StyleWriter.calls, 1)
        self.assertEqual(styled['signals'], [])
        self.assertIn('待人工复核', styled['reason'])

    def test_all_dimensions_at_threshold_do_not_force_rewrite(self):
        self.config['recent_drafts'] = []
        StyleWriter.scores = dict.fromkeys(('title', 'opening', 'hierarchy', 'voice'), 4)
        StyleWriter.scores.update(usefulness=3, distinctness=3)
        out, styled = self.run_pass({'title': '不会采用的改写', 'body': ARTICLE['body']})
        self.assertEqual(out['title'], ARTICLE['title'])
        self.assertIn('未改写', styled['reason'])

    def test_same_facts_unit(self):
        self.assertEqual(same_facts(ARTICLE, dict(ARTICLE)), [])
        self.assertTrue(same_facts(ARTICLE, {**ARTICLE, 'body': ARTICLE['body'] + '另有“新增的一句引语原文”。[S1]'}))


class LibraryTest(unittest.TestCase):
    def setUp(self):
        self.f = test_core.CoreTest()
        self.f.setUp()
        self.api = Api(Backend(self.f.db, test_core.FakeVault(), Path(tempfile.mkdtemp(dir=self.f.path))))

    def tearDown(self):
        self.f.tearDown()

    def test_add_list_delete_and_card(self):
        self.assertEqual(SCHEMA_VERSION, 6)
        listed = self.api.ex_list()
        self.assertTrue(listed['card_is_default'])
        self.assertEqual(listed['examples'], [])
        self.assertTrue(self.api.ex_add('标题', '正文段落。', 'good', '85', '别夸大')['ok'])
        self.assertFalse(self.api.ex_add('标题', '', 'good', '', '')['ok'])
        self.assertFalse(self.api.ex_add('标题', '正文', 'nice', '', '')['ok'])
        rows = self.api.ex_list()['examples']
        self.assertEqual((rows[0]['title'], rows[0]['score'], rows[0]['avoid']), ('标题', 85, '别夸大'))
        self.assertFalse(self.api.style_save({'reader': ''})['ok'])
        card = dict(DEFAULT_STYLE_CARD, voice='像朋友聊天')
        self.assertTrue(self.api.style_save(card)['ok'])
        self.assertEqual(self.api.ex_list()['card']['voice'], '像朋友聊天')
        self.assertTrue(self.api.style_save(None)['ok'])
        self.assertTrue(self.api.ex_list()['card_is_default'])
        self.api.ex_delete(rows[0]['id'])
        self.assertEqual(self.api.ex_list()['examples'], [])

    def test_import_docx_and_markdown(self):
        import docx
        path = self.f.path / '范文.docx'
        document = docx.Document()
        document.add_heading('Word 范文标题', level=1)
        document.add_paragraph('开头一段。')
        document.add_heading('一、小标题', level=2)
        document.add_paragraph('正文。')
        document.save(str(path))
        title, body, meta = read_example_file(path)
        self.assertEqual(title, 'Word 范文标题')
        self.assertIn('## 一、小标题', body)
        with self.assertRaises(ValueError):
            read_example_file(self.f.path / 'x.pdf')

    def test_mark_a_draft_as_example_or_counter_example(self):
        job = self.f.db.job('mark:1', self.f.config)
        draft_id = self.f.db.save_draft(job['id'], '稿件标题', '正文。[S1]\n\n## 小标题\n\n更多。[W1]', [], '报告')
        self.assertTrue(self.api.dr_mark_example(draft_id, 'bad', '太书面', '40')['ok'])
        row = self.f.db.examples()[0]
        self.assertEqual((row['kind'], row['score'], row['avoid'], row['origin']), ('bad', 40, '太书面', f'稿件 #{draft_id}'))
        self.assertNotIn('[S1]', row['body'])

    def test_generated_card_is_validated(self):
        card = dict(DEFAULT_STYLE_CARD)
        transport = SequenceTransport(chat(card), chat({**card, 'taboos': []}))
        model = Model('deepseek', defaults()['providers']['deepseek'], 'k', transport)
        examples = [{'kind': 'good', 'skeleton': '骨架', 'score': 80, 'avoid': ''}, {'kind': 'bad', 'skeleton': '反例', 'avoid': '太书面'}]
        self.assertEqual(model.style_card(examples)['voice'], card['voice'])
        sent = user_json(transport.payloads[0])
        self.assertEqual(len(sent['examples']), 1)
        self.assertEqual(sent['counter_examples'][0]['why'], '太书面')
        with self.assertRaises(ApiError):
            model.style_card(examples)
        with self.assertRaises(ValueError):
            model.style_card([{'kind': 'bad', 'skeleton': 'x'}])

    def test_a_positive_example_never_uses_a_forbidden_wording(self):
        # The conflict seen in the first real card: the title rule praised “缺一样直接驳回”, the taboos banned it.
        card = {**DEFAULT_STYLE_CARD, 'title': '用紧迫感点明利害，如“验收潮来了”“缺一样直接驳回”。',
                'taboos': ['不得使用原文没有的夸大说法，如“缺一样直接驳回”'],
                'examples': [{'fact': '要准备资料', 'good': '缺一样直接驳回！', 'bad': '资料不全会被拉黑'},
                             {'fact': '试运行30天', 'good': '实打实跑满30天', 'bad': '跑不满直接追回补贴'},
                             {'fact': '延期要报备', 'good': '别拖着不报', 'bad': '私自延期标记异常'}]}
        avoid = '原文没有的“私自变更……验收直接不予通过”“别拖着”'
        cleaned, fixes = clean_style_card(card, [avoid])
        self.assertEqual(cleaned['title'], '用紧迫感点明利害，如“验收潮来了”。')
        self.assertEqual([e['good'] for e in cleaned['examples']], ['实打实跑满30天'])   # taboo and avoid list both apply
        self.assertEqual(len(fixes), 3)
        self.assertEqual(cleaned['taboos'], card['taboos'])                            # taboos keep their quotes
        self.assertIn('缺一样直接驳回', card['title'])                                   # the input is not changed
        fragments = clean_style_card({**card, 'title': '短句，如“私自变更后验收直接不予通过”'}, [avoid])[0]
        self.assertEqual(fragments['title'], '短句')                                   # “……” matches in fragments
        self.assertEqual(clean_style_card(DEFAULT_STYLE_CARD)[1], [])                   # a clean card is left alone

    def test_generation_cleans_the_card_and_says_so_in_the_prompt(self):
        card = {**DEFAULT_STYLE_CARD, 'title': '点明利害，如“缺一样直接驳回”。',
                'examples': [{'fact': 'a', 'good': '缺一样直接驳回', 'bad': '编造'}]}
        transport = SequenceTransport(chat(card), chat({**card, 'title': '点明利害。', 'examples': DEFAULT_STYLE_CARD['examples']}))
        model = Model('deepseek', defaults()['providers']['deepseek'], 'k', transport)
        examples = [{'kind': 'good', 'skeleton': '骨架', 'score': 80, 'avoid': '“缺一样直接驳回”'}]
        with self.assertRaises(ApiError):                     # every example was forbidden: not saved
            model.style_card(examples)
        self.assertIn('绝不能出现“不要学”', system_text(transport.payloads[0]))
        self.assertEqual(model.style_card(examples)['title'], '点明利害。')

    def test_choose_example_prefers_the_closest_good_one(self):
        examples = [{'id': 1, 'kind': 'good', 'skeleton': '标题：城市更新入围名单', 'score': 90},
                    {'id': 2, 'kind': 'good', 'skeleton': '标题：设备更新验收清单', 'score': 70},
                    {'id': 3, 'kind': 'bad', 'skeleton': '标题：设备更新验收要求', 'score': 20}]
        self.assertEqual(choose_example(examples, '设备更新验收')['id'], 2)
        self.assertIsNone(choose_example([], '任何'))


class ConfigTest(unittest.TestCase):
    def test_defaults_and_migration(self):
        config = migrate({})
        self.assertIsNone(config['style_card'])
        self.assertTrue(config['style_review_enabled'])


if __name__ == '__main__':
    unittest.main()
