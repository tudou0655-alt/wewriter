"""Desktop UI: five pages (工作台、公众号、文章库、稿件库、设置) on Tkinter/ttk.

Work that touches the network or a model runs on one worker thread; results
come back through ``self.events`` and are applied on the Tk main thread.
"""
import json
import os
from pathlib import Path
import queue
import re
import sqlite3
import sys
import threading
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from tkinter import font as tkfont
from tkinter.scrolledtext import ScrolledText
import uuid
import webbrowser
from datetime import datetime, timedelta, timezone
from zoneinfo import ZoneInfo
from . import BUILD_INFO_NAME, __version__, build_label
from .api import ApiError, Model, Transport, wechat_url
from .config import due_key, validate
from .db import Database
from .engine import Engine, format_coverage, format_funnel, load_funnel
from .logs import logger, redact
from .windows import startup_enabled, set_startup

CHINA = timezone(timedelta(hours=8))

ERROR_LABELS = {'material': '素材不可用', 'no_accounts': '未添加公众号', 'collect': '采集失败',
                'auth': '登录/密钥', 'network': '网络', 'model_output': '模型输出', 'api': '接口错误',
                'config': '设置问题', 'cancelled': '已停止', 'internal': '程序异常'}
ERROR_HINTS = {
    'material': '这篇文章无法提供可用正文（已删除、验证页或图片为主）。批量任务会自动由下一篇补位；也可扩大时间范围后新建任务。',
    'no_accounts': '到“公众号”页扫码同步书架或启用公众号。',
    'collect': '所有公众号都没能读取列表。到“设置 → 服务状态”检查采集服务，必要时重新扫码登录。',
    'auth': '微信读书登录失效或 API Key/余额/权限有问题。微信读书类错误请重新扫码；HTTP 401/402/403 请检查设置中的 Key 和账户余额。',
    'network': '网络连接失败或超时。检查网络后重试，已完成的步骤不会重复计费。',
    'model_output': '模型两次输出都没通过格式或来源校验。直接重试通常即可；反复出现可换写作模型。',
    'api': '接口返回错误，详情见结果说明；可稍后重试。',
    'config': '请按结果说明调整设置后重试。',
    'internal': '程序内部异常，详情已写入数据目录 logs\\wewriter.log。可以先重试，反复出现请把日志发给开发者。',
}

# Model IDs offered in settings; any other ID the service supports can be typed.
# Checked against the providers' model pages on 2026-10-25; the user chose these four.
MODEL_OPTIONS = {'openai': ['gpt-5.6-terra', 'gpt-6-luna'],
                 'deepseek': ['deepseek-v4-pro', 'deepseek-flash']}
RECOMMENDED = {'writer_provider': 'deepseek', 'summary_provider': 'deepseek',
               'deepseek': 'deepseek-v4-pro', 'search_model': 'gpt-5.6-terra'}
STYLE_PRESETS = {
    '政策申报实务': '面向企业决策者和申报负责人：专业但好懂，结论先行，像懂行的同行讲清政策变化和申报要点；用数据和往年对比说话，避免公文套话和夸张营销。',
    '参考原文风格': '参考原文的语气、结构和表达习惯，但不照搬原句。',
    '通俗自然': '面向普通读者，通俗自然，清晰有逻辑，避免夸张营销和空泛套话。',
    '专业简洁': '专业简洁，结论先行，用词准确，避免口语化和空泛套话。',
}
EXPORT_FORMATS = {'docx': 'Word（.docx）', 'md': 'Markdown（.md）', 'txt': '纯文本（.txt）'}
RANGES = (('今天', 'today', 24), ('近 24 小时', 'rolling', 24), ('近 3 天', 'rolling', 72))
PAGES = ('工作台', '公众号', '文章库', '稿件库', '设置')
# Old page names (smoke test, scripts) map onto the five pages.
PAGE_ALIASES = {'概览': '工作台', '运行记录': '工作台', '公众号管理': '公众号', '文章资料库': '文章库',
                '稿件工作台': '稿件库'}
PHASES = ('采集', '预检', '写作', '二审', '待审核')
RETRYABLE = ('失败', '部分完成', '已中断', '已停止', '等待中')

LIGHT = {'paper': '#ffffff', 'base': '#f4f6f8', 'ink': '#172638', 'muted': '#627183', 'line': '#dfe5ec',
         'accent': '#225ac3', 'accent_hover': '#1b4ba6', 'tint': '#edf3ff', 'good': '#206c53',
         'warn': '#8a5b16', 'bad': '#b42318', 'good_bg': '#e7f4ee', 'warn_bg': '#fbf1df', 'bad_bg': '#fdecea',
         'field': '#ffffff', 'select': '#dce8ff', 'on_accent': '#ffffff', 'border': '#c3ccd7'}
DARK = {'paper': '#20252d', 'base': '#14181f', 'ink': '#eef2f7', 'muted': '#aab6c6', 'line': '#38414d',
        'accent': '#98baff', 'accent_hover': '#b5ceff', 'tint': '#243859', 'good': '#8dd3b7',
        'warn': '#edc788', 'bad': '#f4a59b', 'good_bg': '#1f3a31', 'warn_bg': '#3d3220', 'bad_bg': '#45231f',
        'field': '#181c23', 'select': '#2c4470', 'on_accent': '#10151c', 'border': '#4a5463'}
FONT = 'Microsoft YaHei UI'


def windows_dark_mode():
    if os.environ.get('WEWRITER_THEME') in ('light', 'dark'):
        return os.environ['WEWRITER_THEME'] == 'dark'
    if os.name != 'nt':
        return False
    try:
        import winreg
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER,
                            r'Software\Microsoft\Windows\CurrentVersion\Themes\Personalize') as key:
            return winreg.QueryValueEx(key, 'AppsUseLightTheme')[0] == 0
    except OSError:
        return False


def reduce_motion():
    """Honour Windows' “show animations” switch (and an override for tests)."""
    if os.environ.get('WEWRITER_REDUCE_MOTION') == '1':
        return True
    if os.name != 'nt':
        return False
    try:
        import ctypes
        enabled = ctypes.c_int(1)
        # SPI_GETCLIENTAREAANIMATION
        ctypes.windll.user32.SystemParametersInfoW(0x1042, 0, ctypes.byref(enabled), 0)
        return not enabled.value
    except Exception:
        return False


def asset_path(name):
    base = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parents[1]))
    return base / "assets" / name


def packaged_build_info():
    """build-info.json bundled by scripts/build.ps1; None for source runs or unrecorded builds."""
    base = getattr(sys, '_MEIPASS', None)
    if not base:
        return None
    try:
        info = json.loads((Path(base) / BUILD_INFO_NAME).read_text(encoding='utf-8'))
    except (OSError, ValueError):
        return None
    return info if isinstance(info, dict) else None


def version_text():
    return build_label(packaged_build_info(), frozen=bool(getattr(sys, 'frozen', False)))


def job_status_label(job):
    """A finished task whose warnings remain is flagged so a full draft count cannot hide them."""
    try:
        warned = bool(json.loads(job.get('warnings') or '[]'))
    except ValueError:
        warned = True
    if job['status'] == '等待中' and job.get('stage') == '排队中':
        return '排队中'
    return job['status'] + ('（有提示）' if warned and job['status'] == '已完成' else '')


def job_detail_text(job):
    """Run-record detail; every diagnostic is redacted again at the display boundary."""
    hint = ERROR_HINTS.get(job.get('error_code') or '', '')
    try:
        warnings = json.loads(job.get('warnings') or '[]')
    except ValueError:
        warnings = [job.get('warnings')]
    text = (f"任务 #{job['id']}\n状态：{job_status_label(job)}\n阶段：{job['stage']}\n{job['message']}\n"
            + (f"\n处理建议：{hint}\n" if hint else '')
            + "\n采集漏斗：\n" + format_funnel(job.get('funnel'))
            + "\n\n警告：\n" + ("\n".join(str(w) for w in warnings) if warnings else '无')
            + "\n\n“继续原任务”沿用本任务的原配置、时间窗和已选文章，已完成的稿件与步骤不重复；"
              "如需获取新文章或使用当前设置，请点击“按当前配置重新采集”。"
            + "\n\n配置快照（不含密钥）：\n" + json.dumps(json.loads(job["config"]), ensure_ascii=False, indent=2))
    return redact(text)


def china_time(value, pattern="%m-%d %H:%M"):
    """Display time is always China time (UTC+8), whatever the PC's time zone."""
    try:
        when = datetime.fromisoformat(value)
    except (ValueError, TypeError):
        return str(value or '')
    if when.tzinfo is None:
        when = when.replace(tzinfo=timezone.utc)
    return when.astimezone(CHINA).strftime(pattern)


local_time = china_time


def export_document(title, body, sources, markdown=True):
    """Create a reader-facing export without internal S/W citation markers."""
    headings = list(re.finditer(r'(?im)^#{1,6}\s*(?:参考来源|参考资料|资料来源|引用来源)\s*$', body))
    if headings:
        body = body[:headings[-1].start()].rstrip()
    cited = []
    for source_id in re.findall(r'\[([SW]\d+)\]', body):
        if source_id not in cited:
            cited.append(source_id)
    body = re.sub(r'\s*\[(?:S|W)\d+\]', '', body)
    body = re.sub(r'[ \t]+([，。！？；：,.!?;:])', r'\1', body)
    body = re.sub(r'\n{3,}', '\n\n', body).strip()
    by_id = {str(s.get('source_id')): s for s in sources if isinstance(s, dict)}
    selected = [by_id[sid] for sid in cited if sid in by_id]
    if not selected:
        selected = [s for s in sources if isinstance(s, dict) and s.get('url')]
    unique, seen = [], set()
    for source in selected:
        url = str(source.get('url') or '').strip()
        if not url or url in seen:
            continue
        seen.add(url)
        unique.append(source)
    if markdown:
        text = f"# {title.strip()}\n\n{body}"
        if unique:
            text += "\n\n## 参考资料网站\n\n" + "\n".join(
                f"- [{s.get('title') or s.get('account_name') or '资料来源'}]({s['url']})" for s in unique)
    else:
        text = f"{title.strip()}\n\n{body}"
        if unique:
            text += "\n\n参考资料网站\n\n" + "\n".join(
                f"- {s.get('title') or s.get('account_name') or '资料来源'}：{s['url']}" for s in unique)
    return text.strip() + "\n"


def source_website_text(body, sources):
    exported = export_document('', body, sources, markdown=False)
    marker = '\n参考资料网站\n\n'
    return exported.split(marker, 1)[1].strip() if marker in exported else '当前正文没有可显示的参考网站。'


def clipboard_text(title, body, sources):
    """Clean copy for the clipboard, built from the same export so [S1]/[W1] marks never leave.

    Returns (plain, html): plain text for any editor, and HTML that keeps headings, lists, bold and
    tables when pasted into the WeChat editor or Word.
    """
    import html as _html

    def rich(text):
        # Escape quotes too: Markdown URLs are model-controlled text and later become href attributes.
        text = _html.escape(text, quote=True)
        text = re.sub(r'\*\*(.+?)\*\*', r'<strong>\1</strong>', text)
        return re.sub(r'\[([^\]]+)\]\((https?://[^)\s]+)\)', lambda m: f'<a href="{m.group(2)}">{m.group(1)}</a>', text)

    plain, rich_blocks = [], []
    for block in re.split(r'\n{2,}', export_document(title, body, sources, markdown=True).strip()):
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        if not lines:
            continue
        heading = re.match(r'^(#{1,6})\s+(.*)$', lines[0])
        if heading and len(lines) == 1:
            level = min(len(heading.group(1)), 3)
            plain.append(_plain_inline(heading.group(2)))
            rich_blocks.append(f'<h{level}>{rich(heading.group(2))}</h{level}>')
        elif all(line.startswith('>') for line in lines):
            quoted = [re.sub(r'^>\s?', '', line) for line in lines]
            plain.append('\n'.join(_plain_inline(line) for line in quoted))
            rich_blocks.append('<blockquote>' + '<br>'.join(rich(line) for line in quoted) + '</blockquote>')
        elif len(lines) > 1 and all(line.startswith('|') for line in lines):
            rows = [[cell.strip() for cell in line.strip('|').split('|')] for line in lines
                    if not re.fullmatch(r'\|[\s:|-]+\|', line)]
            plain.append('\n'.join('\t'.join(_plain_inline(cell) for cell in row) for row in rows))
            cells = lambda row, tag: ''.join(f'<{tag} style="border:1px solid #ccc;padding:4px 8px">{rich(c)}</{tag}>' for c in row)
            rich_blocks.append('<table style="border-collapse:collapse"><thead><tr>' + cells(rows[0], 'th') + '</tr></thead><tbody>'
                               + ''.join('<tr>' + cells(row, 'td') + '</tr>' for row in rows[1:]) + '</tbody></table>')
        elif all(re.match(r'^[-*]\s+', line) for line in lines):
            items = [re.sub(r'^[-*]\s+', '', line) for line in lines]
            plain.append('\n'.join('- ' + _plain_inline(item) for item in items))
            rich_blocks.append('<ul>' + ''.join(f'<li>{rich(item)}</li>' for item in items) + '</ul>')
        elif all(re.match(r'^\d+[.、]\s*', line) for line in lines):
            items = [re.sub(r'^\d+[.、]\s*', '', line) for line in lines]
            plain.append('\n'.join(f'{i}. {_plain_inline(item)}' for i, item in enumerate(items, 1)))
            rich_blocks.append('<ol>' + ''.join(f'<li>{rich(item)}</li>' for item in items) + '</ol>')
        else:
            plain.append('\n'.join(_plain_inline(re.sub(r'^#{1,6}\s+', '', line)) for line in lines))
            rich_blocks.append('<p>' + '<br>'.join(rich(re.sub(r'^#{1,6}\s+', '', line)) for line in lines) + '</p>')
    return '\n\n'.join(plain) + '\n', '\n'.join(rich_blocks)


def _plain_inline(text):
    text = re.sub(r'\[([^\]]+)\]\((https?://[^)\s]+)\)', r'\1：\2', text)
    return re.sub(r'(\*\*|__|`)', '', text)


def write_docx(path, markdown_text):
    """Word file from the clean Markdown export (headings, lists, quotes, paragraphs)."""
    from docx import Document
    document = Document()
    for line in markdown_text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        heading = re.match(r'^(#{1,6})\s+(.*)$', stripped)
        if heading:
            document.add_heading(_plain_inline(heading.group(2)), level=min(len(heading.group(1)) - 1, 4))
        elif re.match(r'^[-*]\s+', stripped):
            document.add_paragraph(_plain_inline(re.sub(r'^[-*]\s+', '', stripped)), style='List Bullet')
        elif re.match(r'^\d+[.、]\s*', stripped):
            document.add_paragraph(_plain_inline(re.sub(r'^\d+[.、]\s*', '', stripped)), style='List Number')
        elif stripped.startswith('>'):
            document.add_paragraph(_plain_inline(stripped.lstrip('> ').strip()), style='Intense Quote')
        else:
            document.add_paragraph(_plain_inline(stripped))
    document.save(str(path))


def write_export(path, fmt, title, body, sources):
    """Write the clean publication copy; internal S/W markers never reach the file."""
    path = Path(path)
    if fmt == 'docx':
        write_docx(path, export_document(title, body, sources, markdown=True))
    else:
        path.write_text(export_document(title, body, sources, markdown=fmt == 'md'), encoding='utf-8-sig')
    return path


def write_report(path, fmt, title, report):
    """The material report keeps source mapping and review records; separate from the clean copy."""
    path = Path(path)
    text = f"# 素材报告：{title.strip()}\n\n{report.strip()}\n"
    if fmt == 'docx':
        write_docx(path, text)
    else:
        path.write_text(text if fmt == 'md' else re.sub(r'^#\s*', '', text), encoding='utf-8-sig')
    return path


def phase_index(stage):
    stage = str(stage or '')
    if stage in ('完成', '已完成'):
        return 4
    if any(word in stage for word in ('二审', '审核')):
        return 3
    if any(word in stage for word in ('写作', '研究', '联网', '补位')):
        return 2
    if any(word in stage for word in ('正文', '摘要', '素材', '风格', '图片')):
        return 1
    return 0


def draft_hint(report):
    """Short reviewer hint from the machine report; the report itself is never rewritten."""
    report = report or ''
    if '二审失败' in report:
        return '二审失败，需人工复核'
    if re.search(r'需复核|质量提示', report):
        return '有事实或质量提示待核实'
    # Only the engine's own status section means research failed; the words 降级/未使用 can appear in article notes.
    if '## 联网研究状态' in report:
        return '联网资料降级，仅基于原文'
    return ''


def compose_style(preset, extra):
    base = STYLE_PRESETS.get(preset, preset.strip())
    extra = extra.strip()
    return (base + ('；补充要求：' + extra if extra else '')).strip()


class App:
    def __init__(self, root, db, vault, data_dir, demo=False):
        self.root, self.db, self.vault, self.data_dir, self.demo = root, db, vault, Path(data_dir), demo
        self.events = queue.Queue()
        self.engine = Engine(db, vault, lambda msg: self.events.put(("progress", msg)),
                             profile_dir=self.data_dir / 'weread-edge-profile')
        self.busy = False
        self.closed = False
        self.tray = None
        self.page = None
        self.draft_id = None
        self.editor_dirty = False
        self.settings_dirty = False
        self.current_job_id = None
        self.session_queue = []          # queued job ids created in this session
        self.queue_paused = False
        self.page_timers = []
        self.pal = DARK if windows_dark_mode() else LIGHT
        self.still = reduce_motion()
        self.history_open = False
        self.history_page = 0
        self.article_open = None
        self.sync_state = None
        screen_width = self.root.winfo_screenwidth()
        screen_height = self.root.winfo_screenheight()
        self.compact = (screen_width < 1180 or screen_height < 800
                        or os.environ.get('WEWRITER_COMPACT_UI') == '1')
        self.root.title("WeWriter 公众号工作台" + (" · 演示数据" if demo else ""))
        try:
            self.window_icon = tk.PhotoImage(file=str(asset_path("wewriter-cat.png")))
            self.root.iconphoto(True, self.window_icon)
        except (tk.TclError, OSError):
            self.window_icon = None
        width = max(760, min(1240, screen_width - 40))
        height = max(540, min(860, screen_height - 80))
        self.root.geometry(f"{width}x{height}")
        self.root.minsize(min(880, screen_width - 20), min(580, screen_height - 50))
        self.root.configure(bg=self.pal['base'])
        tkfont.nametofont("TkDefaultFont").configure(family=FONT, size=10)
        self.style()
        self.build_shell()
        self.show("工作台")
        self.root.protocol("WM_DELETE_WINDOW", self.close)
        self.root.report_callback_exception = self.report_callback_exception
        self.root.bind_class('TButton', '<Return>', lambda e: e.widget.invoke())
        self.root.bind_class('TCheckbutton', '<Return>', lambda e: e.widget.invoke())
        self.root.after(150, self.poll)
        self.root.after(2000, self.schedule_tick)

    # ------------------------------------------------------------------ style

    def style(self):
        p = self.pal
        style = ttk.Style(self.root)
        style.theme_use("clam")
        size = 9 if self.compact else 10
        style.configure(".", background=p['base'], foreground=p['ink'], font=(FONT, size),
                        bordercolor=p['line'], lightcolor=p['line'], darkcolor=p['line'],
                        troughcolor=p['base'], fieldbackground=p['field'], selectbackground=p['select'],
                        selectforeground=p['ink'], insertcolor=p['ink'])
        style.configure("TFrame", background=p['base'])
        style.configure("Card.TFrame", background=p['paper'])
        style.configure("TLabel", background=p['base'], foreground=p['ink'])
        style.configure("Card.TLabel", background=p['paper'], foreground=p['ink'])
        style.configure("Muted.TLabel", background=p['paper'], foreground=p['muted'], font=(FONT, size - 1))
        style.configure("Base.Muted.TLabel", background=p['base'], foreground=p['muted'], font=(FONT, size - 1))
        style.configure("Title.TLabel", background=p['base'], foreground=p['ink'],
                        font=(FONT, 17 if self.compact else 20, "bold"))
        style.configure("Section.TLabel", background=p['paper'], foreground=p['ink'],
                        font=(FONT, size + 2, "bold"))
        style.configure("Good.TLabel", background=p['paper'], foreground=p['good'], font=(FONT, size, 'bold'))
        style.configure("Warn.TLabel", background=p['paper'], foreground=p['warn'], font=(FONT, size, 'bold'))
        style.configure("Bad.TLabel", background=p['paper'], foreground=p['bad'], font=(FONT, size, 'bold'))
        pad = (12, 6) if self.compact else (15, 7)
        style.configure("TButton", padding=pad, background=p['paper'], foreground=p['ink'],
                        bordercolor=p["border"], relief="solid", borderwidth=1, font=(FONT, size), width=-6)
        style.map("TButton", background=[("disabled", p['base']), ("pressed", p['select']), ("active", p['tint'])],
                  foreground=[("disabled", p['muted']), ("active", p['accent'])],
                  bordercolor=[("focus", p['accent'])])
        style.configure("Primary.TButton", background=p['accent'], foreground=p['on_accent'],
                        bordercolor=p['accent'], font=(FONT, size, 'bold'))
        style.map("Primary.TButton", background=[("disabled", p['line']), ("pressed", p['accent_hover']),
                                                 ("active", p['accent_hover'])],
                  foreground=[("disabled", p['muted']), ("active", p['on_accent'])])
        style.configure("Danger.TButton", foreground=p['bad'])
        style.map("Danger.TButton", background=[("active", p['bad_bg'])], foreground=[("active", p['bad'])])
        style.configure("Seg.TButton", padding=(10, 5), background=p['paper'])
        style.configure("SegOn.TButton", padding=(10, 5), background=p['tint'], foreground=p['accent'],
                        bordercolor=p['accent'], font=(FONT, size, 'bold'))
        style.map("SegOn.TButton", background=[("active", p['tint'])])
        style.configure("Link.TButton", padding=(4, 2), background=p['paper'], foreground=p['accent'],
                        bordercolor=p['paper'])
        style.map("Link.TButton", background=[("active", p['tint'])])
        style.configure("TEntry", padding=6, fieldbackground=p['field'], foreground=p['ink'])
        style.configure("TCombobox", padding=5, fieldbackground=p['field'], foreground=p['ink'],
                        arrowcolor=p['muted'])
        style.map("TCombobox", fieldbackground=[("readonly", p['field'])], foreground=[("readonly", p['ink'])])
        self.root.option_add('*TCombobox*Listbox.background', p['field'])
        self.root.option_add('*TCombobox*Listbox.foreground', p['ink'])
        self.root.option_add('*TCombobox*Listbox.selectBackground', p['select'])
        self.root.option_add('*TCombobox*Listbox.selectForeground', p['ink'])
        style.configure("Treeview", background=p['paper'], fieldbackground=p['paper'], foreground=p['ink'],
                        rowheight=30 if self.compact else 36, borderwidth=0, font=(FONT, size))
        style.configure("Treeview.Heading", background=p['base'], foreground=p['muted'],
                        padding=(8, 6), relief="flat", font=(FONT, size, "bold"))
        style.map("Treeview.Heading", background=[("active", p['tint'])])
        style.map("Treeview", background=[("selected", p['select'])], foreground=[("selected", p['ink'])])
        style.configure("TCheckbutton", background=p['paper'], foreground=p['ink'])
        style.map("TCheckbutton", background=[("active", p['paper'])])
        style.configure("TRadiobutton", background=p['paper'], foreground=p['ink'])
        style.map("TRadiobutton", background=[("active", p['paper'])])
        style.configure("Vertical.TScrollbar", background=p['base'], troughcolor=p['base'], arrowcolor=p['muted'])
        style.configure("Horizontal.TProgressbar", background=p['accent'], troughcolor=p['line'], borderwidth=0)

    # ------------------------------------------------------------------ shell

    def build_shell(self):
        """Prototype layout: a header bar across the top, navigation on the left."""
        p = self.pal
        header = tk.Frame(self.root, bg=p['paper'], highlightbackground=p['line'], highlightthickness=1)
        header.pack(side='top', fill='x')
        tk.Label(header, text="WeWriter", bg=p['paper'], fg=p['ink'],
                 font=(FONT, 15 if self.compact else 16, "bold")).pack(side='left', padx=(20, 10), pady=12)
        tk.Label(header, text="公众号工作台" + (" · 演示数据" if self.demo else ""), bg=p['paper'], fg=p['muted'],
                 font=(FONT, 9)).pack(side='left', pady=12)
        tk.Label(header, text=version_text(), bg=p['paper'], fg=p['muted'], font=(FONT, 8)).pack(
            side='right', padx=20)
        body = tk.Frame(self.root, bg=p['base'])
        body.pack(fill='both', expand=True)
        sidebar_width = 140 if self.compact else 160
        sidebar = tk.Frame(body, bg=p['paper'], width=sidebar_width, highlightbackground=p['line'],
                           highlightthickness=1)
        sidebar.pack(side="left", fill="y")
        sidebar.pack_propagate(False)
        tk.Frame(sidebar, bg=p['paper'], height=14).pack()
        self.nav = {}
        for title in PAGES:
            button = tk.Button(sidebar, text=title, anchor="w", relief="flat", bd=0, padx=12, pady=10,
                               bg=p['paper'], fg=p['muted'], activebackground=p['tint'],
                               activeforeground=p['accent'], font=(FONT, 10), cursor='hand2', takefocus=True,
                               command=lambda page=title: self.show(page))
            button.pack(fill="x", padx=12, pady=3)
            button.bind('<Enter>', lambda e, b=button, t=title: b.configure(bg=p['tint']) if t != self.page else None)
            button.bind('<Leave>', lambda e, b=button, t=title: b.configure(bg=p['paper']) if t != self.page else None)
            button.bind('<Return>', lambda e, page=title: self.show(page))
            self.nav[title] = button
        right = tk.Frame(body, bg=p['base'])
        right.pack(side="left", fill="both", expand=True)
        footer = ttk.Frame(right, padding=(26, 5))
        footer.pack(fill="x", side="bottom")
        self.status = tk.StringVar(value="准备就绪。" if not self.demo else "演示模式：仅本地样例，不调用 API。")
        ttk.Label(footer, textvariable=self.status, style="Base.Muted.TLabel", wraplength=760).pack(side="left")
        self.progress = ttk.Progressbar(footer, mode="indeterminate", length=110)
        # Every page lives in a scrollable column so small screens and 150 %
        # scaling never hide the actions.
        holder = tk.Frame(right, bg=p['base'])
        holder.pack(fill='both', expand=True)
        self.canvas = tk.Canvas(holder, bg=p['base'], highlightthickness=0, bd=0)
        bar = ttk.Scrollbar(holder, orient='vertical', command=self.canvas.yview)
        self.canvas.configure(yscrollcommand=bar.set)
        bar.pack(side='right', fill='y')
        self.canvas.pack(side='left', fill='both', expand=True)
        self.content = ttk.Frame(self.canvas, padding=(18, 14, 18, 14) if self.compact else (26, 20, 26, 18))
        window = self.canvas.create_window((0, 0), window=self.content, anchor='nw')
        self.content.bind('<Configure>', lambda e: self.canvas.configure(scrollregion=self.canvas.bbox('all')))
        self.canvas.bind('<Configure>', lambda e: self.canvas.itemconfigure(window, width=e.width))
        self.root.bind_all('<MouseWheel>', self._wheel, add='+')

    def _wheel(self, event):
        widget = event.widget
        # Lists and text boxes scroll themselves.
        if isinstance(widget, (tk.Text, ttk.Treeview, tk.Listbox)) or not str(widget).startswith(str(self.root)):
            return
        try:
            if widget.winfo_toplevel() is not self.root:
                return
        except (tk.TclError, AttributeError):
            return
        box = self.canvas.bbox('all')
        if box and box[3] - box[1] > self.canvas.winfo_height():
            self.canvas.yview_scroll(int(-event.delta / 120), 'units')

    def show(self, page):
        page = PAGE_ALIASES.get(page, page)
        if self.editor_dirty and not messagebox.askyesno("未保存修改", "离开后将丢弃未保存的稿件修改。是否离开？", parent=self.root):
            return
        if self.settings_dirty and not messagebox.askyesno("未保存设置", "设置有未保存的修改，离开后将丢弃。是否离开？", parent=self.root):
            return
        self.editor_dirty = self.settings_dirty = False
        self.cancel_page_timers()
        if page != '稿件库':
            self.draft_id = None
        if page != '文章库':
            self.article_open = None
        if page != '公众号':
            self.sync_state = None
        self.page = page
        for child in self.content.winfo_children():
            child.destroy()
        for name, button in self.nav.items():
            active = name == page
            button.configure(bg=self.pal['tint'] if active else self.pal['paper'],
                             fg=self.pal['accent'] if active else self.pal['muted'],
                             font=(FONT, 10, 'bold' if active else 'normal'))
        self.canvas.yview_moveto(0)
        {"工作台": self.workbench, "公众号": self.accounts, "文章库": self.articles,
         "稿件库": self.drafts, "设置": self.settings}[page]()

    def refresh_page(self):
        if self.page and not (self.editor_dirty or self.settings_dirty):
            self.show(self.page)

    def page_after(self, delay, callback):
        timer = self.root.after(delay, callback)
        self.page_timers.append(timer)
        return timer

    def cancel_page_timers(self):
        for timer in self.page_timers:
            try:
                self.root.after_cancel(timer)
            except tk.TclError:
                pass
        self.page_timers = []

    # --------------------------------------------------------------- widgets

    def heading(self, title, *actions):
        bar = ttk.Frame(self.content)
        bar.pack(fill='x', pady=(0, 14))
        ttk.Label(bar, text=title, style="Title.TLabel").pack(side='left')
        for text, action, kind in reversed(actions):
            self.button(bar, text, action, kind).pack(side='right', padx=(8, 0))
        return bar

    def card(self, parent=None, title=None, side=None, fill='x', expand=False, pady=(0, 14)):
        frame = tk.Frame(parent or self.content, bg=self.pal['paper'], highlightbackground=self.pal['line'],
                         highlightthickness=1, padx=16 if self.compact else 20, pady=14 if self.compact else 16)
        frame.pack(fill=fill, expand=expand, pady=pady, side=side or 'top')
        if title:
            head = tk.Frame(frame, bg=self.pal['paper'])
            head.pack(fill='x', pady=(0, 10))
            ttk.Label(head, text=title, style='Section.TLabel').pack(side='left')
            frame.head = head
        return frame

    def bg_of(self, widget):
        try:
            return widget.cget('bg')
        except tk.TclError:      # ttk widgets take their colour from the style
            return self.pal['base']

    def row(self, parent, pady=(0, 6), bg=None):
        frame = tk.Frame(parent, bg=bg or self.pal['paper'])
        frame.pack(fill='x', pady=pady)
        return frame

    def rule(self, parent, pady=0):
        tk.Frame(parent, bg=self.pal['line'], height=1).pack(fill='x', pady=pady)

    def label(self, parent, text='', style='Card.TLabel', textvariable=None, **pack):
        widget = ttk.Label(parent, text=text, style=style, textvariable=textvariable,
                           wraplength=pack.pop('wrap', 560 if self.compact else 820))
        widget.pack(**({'anchor': 'w'} | pack))
        return widget

    def text(self, parent, text, color='ink', size=None, bold=False, **pack):
        """Plain coloured text on a card (meta lines, counts)."""
        widget = tk.Label(parent, text=text, bg=self.bg_of(parent), fg=self.pal[color], justify='left',
                          font=(FONT, size or (9 if self.compact else 10), 'bold' if bold else 'normal'),
                          wraplength=pack.pop('wrap', 0))
        widget.pack(**({'anchor': 'w'} | pack))
        return widget

    def button(self, parent, text, action, kind='secondary'):
        style = {'primary': 'Primary.TButton', 'danger': 'Danger.TButton', 'link': 'Link.TButton'}.get(kind, 'TButton')
        return ttk.Button(parent, text=text, style=style, command=lambda: self.safe(action), cursor='hand2')

    def pack_button(self, parent, text, action, kind='secondary', side='left'):
        button = self.button(parent, text, action, kind)
        button.pack(side=side, padx=(0, 8) if side == 'left' else (8, 0), pady=2)
        return button

    def badge(self, parent, text, tone='muted', side='right'):
        colors = {'good': ('good_bg', 'good'), 'warn': ('warn_bg', 'warn'), 'bad': ('bad_bg', 'bad'),
                  'accent': ('tint', 'accent'), 'muted': ('base', 'muted')}[tone]
        widget = tk.Label(parent, text=text, bg=self.pal[colors[0]], fg=self.pal[colors[1]], padx=9, pady=2,
                          font=(FONT, 9, 'bold'))
        widget.pack(side=side, padx=(6, 0))
        return widget

    def field(self, parent, label, widget_factory, column=None):
        """Label above a control, as in the prototypes; optional grid column for two-up rows."""
        box = tk.Frame(parent, bg=self.pal['paper'])
        if column is None:
            box.pack(fill='x', pady=(0, 10))
        else:
            box.grid(row=0, column=column, sticky='ew', padx=(0, 16) if column == 0 else (0, 0), pady=(0, 10))
        tk.Label(box, text=label, bg=self.pal['paper'], fg=self.pal['ink'], font=(FONT, 9, 'bold')).pack(
            anchor='w', pady=(0, 4))
        widget = widget_factory(box)
        widget.pack(fill='x')
        return widget

    def switch(self, parent, on, command):
        """On/off control drawn like the prototype's toggle."""
        p = self.pal
        canvas = tk.Canvas(parent, width=34, height=20, bg=self.bg_of(parent), highlightthickness=0, cursor='hand2',
                           takefocus=True)

        def draw(state):
            canvas.delete('all')
            fill = p['accent'] if state else p['line']
            canvas.create_oval(1, 1, 19, 19, fill=fill, outline=fill)
            canvas.create_oval(15, 1, 33, 19, fill=fill, outline=fill)
            canvas.create_rectangle(10, 1, 24, 19, fill=fill, outline=fill)
            x = 16 if state else 2
            canvas.create_oval(x + 1, 3, x + 15, 17, fill=p['paper'] if state else p['muted'], outline='')
        draw(on)
        canvas.bind('<Button-1>', lambda e: self.safe(command))
        canvas.bind('<space>', lambda e: self.safe(command))
        return canvas

    def toggle(self, parent, text, opened=False, on_open=None, on_flip=None):
        """Collapsible area (▸ / ▾): the body is built lazily on first open."""
        holder = tk.Frame(parent, bg=self.pal['paper'])
        holder.pack(fill='x', pady=(6, 0))
        body = tk.Frame(holder, bg=self.pal['paper'])
        state = {'open': False, 'built': False}
        button = ttk.Button(holder, style='Link.TButton', cursor='hand2')

        def flip():
            state['open'] = not state['open']
            button.configure(text=('▾ ' if state['open'] else '▸ ') + text)
            if state['open']:
                if not state['built'] and on_open:
                    on_open(body)
                    state['built'] = True
                body.pack(fill='x', pady=(6, 0))
            else:
                body.pack_forget()
            if on_flip:
                on_flip(state['open'])
        button.configure(command=lambda: self.safe(flip), text='▸ ' + text)
        button.pack(anchor='w')
        if opened:
            flip()
        return body

    def list_rows(self, parent, rows, render, empty):
        """Prototype list: one card with a rule between rows; ``render(row_frame, item)``."""
        box = self.card(parent, pady=(0, 10))
        if not rows:
            self.text(box, empty, 'muted')
            return box
        for index, item in enumerate(rows):
            if index:
                self.rule(box, pady=10)
            line = tk.Frame(box, bg=self.pal['paper'])
            line.pack(fill='x')
            render(line, item)
        return box

    def pager(self, parent, state, total, size, refresh):
        bar = self.row(parent, pady=(0, 6), bg=self.bg_of(parent))
        pages = max(1, (total + size - 1) // size)
        state['page'] = max(0, min(state.get('page', 0), pages - 1))
        if pages <= 1:
            return bar
        self.text(bar, f"第 {state['page'] + 1}/{pages} 页", 'muted', side='left')
        nxt = self.pack_button(bar, '下一页', lambda: (state.update(page=state['page'] + 1), refresh()), side='right')
        prv = self.pack_button(bar, '上一页', lambda: (state.update(page=state['page'] - 1), refresh()), side='right')
        if state['page'] <= 0:
            prv.state(['disabled'])
        if state['page'] >= pages - 1:
            nxt.state(['disabled'])
        return bar

    def selection(self, tree):
        selected = tree.selection()
        if not selected:
            raise ValueError("请先选择一条记录")
        return int(selected[0])

    def safe(self, action):
        try:
            return action()
        except sqlite3.IntegrityError:
            messagebox.showerror("无法保存", "该公众号已在名单中，请勿重复添加。", parent=self.root)
        except Exception as exc:
            # UI validation errors are local, network adapters never include raw server responses.
            if not isinstance(exc, (ValueError, ApiError)):
                logger.exception("界面操作异常")
            messagebox.showerror("操作未完成", redact(str(exc)) or type(exc).__name__, parent=self.root)

    def require_idle(self):
        if self.busy:
            raise ValueError("后台正在执行操作，请等待完成或停止任务后再试")

    def require_live(self):
        if self.demo:
            raise ValueError("这是演示模式，请打开正常版配置 API 后使用")

    def text_window(self, title, text, images=(), links=()):
        win = tk.Toplevel(self.root)
        win.title(title)
        win.configure(bg=self.pal['base'])
        width = min(880, self.root.winfo_screenwidth() - 60)
        height = min(680, self.root.winfo_screenheight() - 100)
        win.geometry(f"{width}x{height}")
        if links:
            bar = tk.Frame(win, bg=self.pal['base'], padx=12, pady=8)
            bar.pack(fill='x')
            for label, action in links:
                ttk.Button(bar, text=label, command=lambda a=action: self.safe(a)).pack(side='left', padx=(0, 8))
        editor = ScrolledText(win, wrap="word", padx=20, pady=18, font=(FONT, 11), relief="flat",
                              bg=self.pal['paper'], fg=self.pal['ink'], insertbackground=self.pal['ink'])
        editor.pack(fill="both", expand=True)
        self.fill_text(editor, text, images, width)
        editor.configure(state="disabled")
        return win

    def fill_text(self, editor, text, images=(), width=800):
        editor.insert("1.0", text)
        for n, match in enumerate(re.finditer(r'https?://[^\s<>"）]+', text)):
            url = match.group().rstrip('。，;')
            tag = 'source_url_' + str(n)
            editor.tag_add(tag, f'1.0+{match.start()}c', f'1.0+{match.start() + len(url)}c')
            editor.tag_config(tag, foreground=self.pal['accent'], underline=True)
            editor.tag_bind(tag, '<Button-1>', lambda event, link=url: self.open_link(link))
        photos = []
        if images:
            from PIL import Image, ImageTk
            root = self.data_dir.resolve()
            for item in images:
                try:
                    path = (self.data_dir / item['local_path']).resolve()
                    if not path.is_relative_to(root) or not path.is_file():
                        continue
                    with Image.open(path) as original:
                        picture = original.convert('RGBA')
                    picture.thumbnail((max(360, width - 90), 520), Image.Resampling.LANCZOS)
                    photo = ImageTk.PhotoImage(picture)
                    photos.append(photo)
                    label = f"\n\n—— 原文图片 {item['position']} ——"
                    if item.get('alt_text'):
                        label += '\n' + item['alt_text']
                    editor.insert('end', label + '\n')
                    editor.image_create('end', image=photo)
                except (OSError, ValueError, tk.TclError):
                    continue
        editor._article_photos = photos

    @staticmethod
    def open_link(url):
        if re.match(r'^https://', str(url)):
            webbrowser.open(url)

    def combo(self, parent, var, values, width=12, readonly=True):
        return ttk.Combobox(parent, textvariable=var, values=values, width=width,
                            state='readonly' if readonly else 'normal')

    # -------------------------------------------------------------- workbench

    def enabled_accounts(self):
        return len(self.db.accounts(enabled=True))

    def latest_job(self):
        if self.current_job_id:
            rows = self.db.rows('SELECT * FROM jobs WHERE id=?', (self.current_job_id,))
            if rows:
                return rows[0]
        _, rows = self.db.top_jobs(1)
        return rows[0] if rows else None

    def queued_jobs(self):
        return self.db.rows("SELECT * FROM jobs WHERE status='等待中' AND stage='排队中' ORDER BY id")

    def workbench(self):
        p = self.pal
        config = self.db.config()
        self.heading('工作台')
        new = self.card(title='新建任务')
        mode = config.get('lookback_mode', 'rolling')
        hours = config.get('lookback_hours', 24)
        ranges = list(RANGES)
        chosen = next((label for label, m, h in ranges if m == mode and (m == 'today' or h == hours)), None)
        if chosen is None:
            chosen = f'近 {hours} 小时'
            ranges.append((chosen, 'rolling', hours))
        line = tk.Frame(new, bg=p['paper'])
        line.pack(fill='x')
        range_var = tk.StringVar(value=chosen)
        count_var = tk.StringVar(value=str(config.get('draft_count', 3)))
        for column, (label, var, values, width) in enumerate((
                ('文章范围', range_var, [r[0] for r in ranges], 10),
                ('生成篇数', count_var, ['1', '2', '3', '4', '5'], 3))):
            box = tk.Frame(line, bg=p['paper'])
            box.pack(side='left', padx=(0, 14))
            tk.Label(box, text=label, bg=p['paper'], fg=p['muted'], font=(FONT, 9)).pack(anchor='w', pady=(0, 4))
            self.combo(box, var, values, width).pack(anchor='w')
        accounts = self.enabled_accounts()
        self.text(line, f'{accounts} 个公众号 · 热度优先' if accounts else '尚未启用公众号（到“公众号”页添加）',
                  'muted', side='left', anchor='s', padx=(0, 14), pady=(0, 6))

        def start():
            self.require_live()
            if not self.enabled_accounts():
                raise ValueError('请先在“公众号”页添加并启用至少一个公众号')
            latest = self.db.config()
            _, mode_value, hours_value = next(r for r in ranges if r[0] == range_var.get())
            latest.update(draft_count=int(count_var.get()), lookback_mode=mode_value, lookback_hours=hours_value,
                          writing_mode='separate_web', selection_mode='hot', comments_enabled=False)
            validate(latest)
            # The choices are also remembered as the defaults of the next task.
            self.db.save_config(latest)
            self.start_task(latest)
        start_button = self.button(line, '加入队列' if self.busy else '开始任务', start, 'primary')
        start_button.pack(side='left', anchor='s')
        queued = self.queued_jobs()
        if queued:
            self.text(line, f'排队中 {len(queued)} 个', 'warn', side='left', anchor='s', padx=(12, 0), pady=(0, 6))

        current = self.card(title='当前任务')
        self.current_card(current)
        drafts_holder = tk.Frame(self.content, bg=p['base'])
        drafts_holder.pack(fill='x')
        self.drafts_card(drafts_holder)
        self.history_holder = tk.Frame(self.content, bg=p['base'])
        self.history_holder.pack(fill='x')
        if self.history_open:
            self.history_card(self.history_holder)
        self.live = {'current': current, 'drafts': drafts_holder}
        if self.busy:
            self.page_after(1500, self.refresh_workbench)

    def refresh_workbench(self):
        """While a task runs, redraw only the live cards from the stored task events."""
        if self.page != '工作台':
            return
        live = getattr(self, 'live', {})
        current, holder = live.get('current'), live.get('drafts')
        if not current or not current.winfo_exists():
            return
        self.cancel_page_timers()
        for child in list(current.winfo_children()):
            if child is not getattr(current, 'head', None):
                child.destroy()
        for child in list(current.head.winfo_children())[1:]:
            child.destroy()
        for child in holder.winfo_children():
            child.destroy()
        self.current_card(current)
        self.drafts_card(holder)
        if self.busy:
            self.page_after(1500, self.refresh_workbench)

    def flip_history(self):
        self.history_open = not self.history_open
        for child in self.history_holder.winfo_children():
            child.destroy()
        if self.history_open:
            self.history_card(self.history_holder)

    def current_card(self, card):
        p = self.pal
        head = card.head
        job = self.latest_job()
        queued = self.queued_jobs()
        self.pack_button(head, '收起历史' if self.history_open else '历史任务', self.flip_history, side='right')
        if job is None:
            self.text(card, '还没有任务。选好文章范围和篇数后点“开始任务”。', 'muted')
            return
        try:
            config = json.loads(job['config'])
        except ValueError:
            config = {}
        target = config.get('draft_count', 3)
        running = self.busy and job['id'] == self.current_job_id
        status = '生成中' if running else job_status_label(job)
        tone = ('good' if running or status == '已完成' else 'bad' if status in ('失败', '已中断') else 'warn')
        self.badge(head, status, tone)
        drafts = self.db.task_drafts(job['id'])
        states = [Database.review_state(d['status']) for d in drafts]
        waiting = states.count('待审核')
        passed = states.count('已通过')
        processing = max(0, target - len(drafts)) if running else 0
        parts = []
        if passed:
            parts.append(f'{passed} 篇已通过')
        if waiting:
            parts.append(f'{waiting} 篇待审核')
        if states.count('需修改'):
            parts.append(f"{states.count('需修改')} 篇需修改")
        if processing:
            parts.append(f'{processing} 篇处理中')
        if not parts:
            parts.append('排队中' if status == '排队中' else '尚无完成的稿件')
        self.text(card, '，'.join(parts), 'ink', 15 if self.compact else 16, True, anchor='w')
        try:
            end = datetime.now(timezone.utc) if running else datetime.fromisoformat(job['updated_at'])
            seconds = max(0, int((end - datetime.fromisoformat(job['created_at'])).total_seconds()))
            spent = (f"{seconds // 3600} 小时 {seconds % 3600 // 60} 分" if seconds >= 3600 else
                     f"{seconds // 60} 分 {seconds % 60} 秒")
        except (ValueError, TypeError):
            spent = '—'
        span = {'today': '今天'}.get(config.get('lookback_mode'), f"近 {config.get('lookback_hours', 24)} 小时")
        self.text(card, f"{span} · 目标 {target} 篇 · 已用时 {spent} · {china_time(job['created_at'])} 开始",
                  'muted', anchor='w', pady=(2, 10))
        # Phase line: completed phases ✓, the running one shimmers, the rest are numbered.
        index = phase_index(job['stage'])
        if status == '已完成' or (not running and drafts and job['status'] in ('已完成', '部分完成')):
            index = len(PHASES) - 1 if waiting or states.count('需修改') else len(PHASES)
        flow = tk.Canvas(card, height=40, bg=p['paper'], highlightthickness=0, cursor='hand2')
        flow.pack(fill='x')
        flow.bind('<Button-1>', lambda e: self.safe(lambda: self.job_detail(job['id'])))
        shimmer = {'offset': 0}

        def draw():
            flow.delete('all')
            width = max(flow.winfo_width(), 400)
            step = width / len(PHASES)
            for i, name in enumerate(PHASES):
                x0, x1 = i * step, (i + 1) * step - 8
                done = i < index
                now = i == index and running
                color = p['good'] if done else p['accent'] if now else p['line']
                flow.create_rectangle(x0, 2, x1, 5, fill=color, width=0)
                if now and not self.still:
                    band = x0 + (shimmer['offset'] % 100) / 100 * (x1 - x0)
                    flow.create_rectangle(max(x0, band - 30), 2, min(x1, band + 30), 5, fill=p['tint'], width=0)
                label = ('✓ ' if done else f'{i + 1} ') + name
                flow.create_text(x0, 22, text=label, anchor='w', font=(FONT, 9, 'bold' if now else 'normal'),
                                 fill=p['accent'] if now else p['ink'] if done else p['muted'])

        def animate():
            if self.page != '工作台' or not running or not flow.winfo_exists():
                return
            shimmer['offset'] += 3
            draw()
            self.page_after(50, animate)
        flow.bind('<Configure>', lambda e: draw())
        if running and not self.still:
            self.page_after(50, animate)
        if job['message'] and not running:
            self.text(card, redact(job['message'])[:200], 'muted', anchor='w', pady=(4, 0), wrap=900)
        actions = self.row(card, pady=(12, 0))
        self.pack_button(actions, '查看详情', lambda: self.job_detail(job['id']))
        if running:
            self.pack_button(actions, '停止任务', self.stop_task)
        elif job['status'] in RETRYABLE and status != '排队中':
            self.pack_button(actions, '继续原任务', lambda: self.resume_job(job['id']))
        if queued:
            restored = [q for q in queued if q['id'] not in self.session_queue]
            self.text(actions, f"排队 {len(queued)} 个" + ('（含上次退出前未开始的任务，需手动开始）' if restored else ''),
                      'warn', side='left', padx=(8, 0))
            self.pack_button(actions, '取消排队', self.cancel_queue, side='right')
            if restored and not self.busy:
                self.pack_button(actions, '开始排队任务', self.start_restored_queue, 'primary', side='right')

    def drafts_card(self, parent):
        job = self.latest_job()
        card = self.card(parent, title='本次稿件')
        self.rule(card)
        if job is None:
            self.text(card, '任务生成的稿件会在这里逐篇出现，单篇完成即可查看。', 'muted', anchor='w', pady=(10, 0))
            return
        try:
            target = json.loads(job['config']).get('draft_count', 3)
        except ValueError:
            target = 0
        drafts = self.db.task_drafts(job['id'])
        running = self.busy and job['id'] == self.current_job_id
        slots = [(d['title'], Database.review_state(d['status']), d['id']) for d in drafts]
        if running:
            slots += [('（写作中）', '处理中', None)] * max(0, target - len(drafts))
        if not slots:
            self.text(card, '本任务还没有完成的稿件。', 'muted', anchor='w', pady=(10, 0))
            return
        for number, (title, state, draft_id) in enumerate(slots, 1):
            if number > 1:
                self.rule(card)
            line = tk.Frame(card, bg=self.pal['paper'])
            line.pack(fill='x', pady=10)
            self.text(line, f'{number:02d}', 'muted', side='left', padx=(0, 22))
            info = tk.Frame(line, bg=self.pal['paper'])
            info.pack(side='left', fill='x', expand=True)
            self.text(info, title, 'ink')
            self.text(info, state, {'已通过': 'good', '需修改': 'warn'}.get(state, 'muted'), 9)
            button = self.button(line, '查看稿件', lambda d=draft_id: self.open_draft(d))
            button.pack(side='right')
            if draft_id is None:
                button.state(['disabled'])

    def history_card(self, parent):
        card = self.card(parent, title='历史任务')
        self.text(card.head, '最近 10 条 · 中国时间', 'muted', side='right')
        self.rule(card)
        holder = tk.Frame(card, bg=self.pal['paper'])
        holder.pack(fill='x')
        state = {'page': self.history_page}

        def render():
            self.history_page = state['page']
            for child in holder.winfo_children():
                child.destroy()
            total, rows = self.db.top_jobs(10, state['page'] * 10)
            if not rows:
                self.text(holder, '暂无历史任务。', 'muted', anchor='w', pady=(10, 0))
                return
            for number, job in enumerate(rows):
                if number:
                    self.rule(holder)
                try:
                    target = json.loads(job['config']).get('draft_count', 3)
                except ValueError:
                    target = '—'
                line = tk.Frame(holder, bg=self.pal['paper'])
                line.pack(fill='x', pady=8)
                line.columnconfigure(0, weight=1, uniform='h')
                line.columnconfigure(1, weight=1, uniform='h')
                tk.Label(line, text=f"{china_time(job['created_at'])} · {job_status_label(job)}", bg=self.pal['paper'],
                         fg=self.pal['ink'], font=(FONT, 10)).grid(row=0, column=0, sticky='w')
                tk.Label(line, text=f"合格稿件 {self.db.qualified_drafts(job['id'])} / {target}", bg=self.pal['paper'],
                         fg=self.pal['muted'], font=(FONT, 9)).grid(row=0, column=1, sticky='w')
                self.button(line, '详情', lambda j=job['id']: self.job_detail(j)).grid(row=0, column=2, sticky='e')
            self.pager(holder, state, total, 10, render)
        render()

    def job_detail(self, job_id):
        job = self.db.rows("SELECT * FROM jobs WHERE id=?", (job_id,))[0]
        links = []
        if job['status'] in RETRYABLE and not (job['status'] == '等待中' and job['stage'] == '排队中'):
            links.append(('继续原任务（沿用原配置）', lambda: self.resume_job(job_id)))
        links.append(('按当前配置重新采集', self.collect_again))
        self.text_window('任务详情', job_detail_text(job), links=links)

    def resume_job(self, job_id):
        self.require_live()
        self.require_idle()
        job = self.db.rows("SELECT * FROM jobs WHERE id=?", (job_id,))[0]
        if job["status"] not in RETRYABLE:
            raise ValueError("该任务已结束，无需继续；如需获取新文章，请点击“按当前配置重新采集”")
        self.launch(job, "正在继续原任务（沿用原配置、时间窗和已选文章）…")

    def collect_again(self):
        self.require_live()
        if not messagebox.askyesno(
                "按当前配置重新采集",
                "将使用当前设置新建任务，时间范围从现在起计算，并重新读取公众号文章列表。\n"
                "原任务及其稿件不会被修改；已用于稿件的文章仍会被排除。\n"
                "新任务会调用采集服务和模型接口，可能产生费用。是否继续？", parent=self.root):
            return
        self.start_task(self.db.config())

    def start_task(self, config):
        """Start now, or queue behind the running work with a frozen config snapshot."""
        key = 'manual:' + uuid.uuid4().hex
        job = self.db.job(key, config)
        if self.busy:
            self.db.update_job(job['id'], stage='排队中', message='排队中：等待当前任务完成后自动开始')
            self.session_queue.append(job['id'])
            self.queue_paused = False
            self.status.set(f"已加入队列（排队 {len(self.queued_jobs())} 个）")
            self.refresh_page()
            return
        self.queue_paused = False
        self.launch(job, "任务已开始：采集文章…")

    def launch(self, job, status):
        config = json.loads(job['config'])
        if job['stage'] == '排队中':
            self.db.update_job(job['id'], stage='采集', message='')
        self.current_job_id = job['id']
        self.start_worker(lambda: self.engine.run(job['run_key'], config), status)
        self.refresh_page()

    def next_queued(self):
        if self.busy or self.queue_paused:
            return
        for job in self.queued_jobs():
            if job['id'] in self.session_queue:
                self.session_queue.remove(job['id'])
                self.launch(job, "开始排队任务…")
                return

    def start_restored_queue(self):
        self.require_live()
        self.require_idle()
        queued = self.queued_jobs()
        if not queued:
            return
        if not messagebox.askyesno('开始排队任务', f'将依次执行 {len(queued)} 个排队任务，会调用采集服务和模型接口，'
                                   '可能产生费用。是否开始？', parent=self.root):
            return
        self.session_queue.extend(q['id'] for q in queued if q['id'] not in self.session_queue)
        self.queue_paused = False
        self.next_queued()

    def cancel_queue(self):
        queued = self.queued_jobs()
        if not queued:
            return
        if not messagebox.askyesno('取消排队', f'取消 {len(queued)} 个尚未开始的任务？正在运行的任务不受影响。',
                                   parent=self.root):
            return
        for job in queued:
            self.db.update_job(job['id'], status='已取消', stage='已取消', message='排队任务已取消，未开始执行')
            if job['id'] in self.session_queue:
                self.session_queue.remove(job['id'])
        self.refresh_page()

    def stop_task(self):
        if not self.busy:
            return
        if not messagebox.askyesno('停止任务', '停止后保留已有稿件。已经发出的请求会在返回后结束，'
                                   '不能保证立即中断。排队任务不会自动开始。是否停止？', parent=self.root):
            return
        self.queue_paused = True
        self.engine.stop()
        self.status.set('正在停止：当前请求结束后停止')

    # --------------------------------------------------------------- accounts

    def accounts(self):
        self.heading('公众号', ('手动添加', self.add_account_dialog, 'secondary'),
                     ('扫码同步书架', self.open_sync, 'primary'))
        self.sync_holder = tk.Frame(self.content, bg=self.pal['base'])
        self.sync_holder.pack(fill='x')
        if getattr(self, 'sync_state', None):
            self.render_sync()
        bar = ttk.Frame(self.content)
        bar.pack(fill='x', pady=(0, 10))
        query = tk.StringVar()
        entry = ttk.Entry(bar, textvariable=query, width=30)
        entry.pack(side='left')
        self.placeholder(entry, query, '搜索已添加的公众号')
        all_accounts = self.db.accounts()
        self.text(bar, f"{len(all_accounts)} 个公众号 · {sum(a['enabled'] for a in all_accounts)} 个已启用",
                  'muted', side='right')
        holder = tk.Frame(self.content, bg=self.pal['base'])
        holder.pack(fill='x')
        details = self.account_details()

        def toggle(account):
            self.require_idle()
            self.db.execute("UPDATE accounts SET enabled=1-enabled WHERE id=?", (account['id'],))
            self.status.set(f"{account['name']} 已{'停用' if account['enabled'] else '启用'}，从下一个新任务起生效")
            refresh()

        def render(line, account):
            info = details.get(account['identifier'], {})
            left = tk.Frame(line, bg=self.pal['paper'])
            left.pack(side='left', fill='x', expand=True)
            self.text(left, account['name'], 'ink', bold=True)
            if not account['enabled']:
                state, color = '已暂停采集', 'muted'
            elif info.get('failed'):
                state, color = '上次采集失败', 'warn'
            elif info:
                state, color = info['state'], 'muted'
            else:
                state, color = '尚未采集', 'muted'
            self.text(left, state, color, 9, pady=(2, 0), anchor='w')
            self.pack_button(line, '详情', lambda a=account: self.account_detail(a, details), side='right')
            tk.Label(line, text='已启用' if account['enabled'] else '已停用', bg=self.pal['paper'], fg=self.pal['ink'],
                     font=(FONT, 9)).pack(side='right', padx=(6, 22))
            self.switch(line, bool(account['enabled']), lambda a=account: toggle(a)).pack(side='right')

        def refresh(*_):
            for child in holder.winfo_children():
                child.destroy()
            text = query.get().strip()
            if text == '搜索已添加的公众号':
                text = ''
            rows = [a for a in self.db.accounts() if not text or text in a['name']]
            self.list_rows(holder, rows, render, '没有匹配的公众号' if text else '还没有公众号，点“扫码同步书架”导入。')
        query.trace_add('write', refresh)
        refresh()

    def placeholder(self, entry, var, hint):
        """Grey hint text inside an empty entry (Tk has no native placeholder)."""
        muted, ink = self.pal['muted'], self.pal['ink']

        def show_hint(*_):
            if not var.get():
                entry.configure(foreground=muted)
                var.set(hint)

        def clear_hint(*_):
            if var.get() == hint:
                var.set('')
                entry.configure(foreground=ink)
        entry.bind('<FocusIn>', clear_hint)
        entry.bind('<FocusOut>', show_hint)
        show_hint()

    def account_details(self):
        """Per-account result of the latest task that collected it (verifiable facts only)."""
        found = {}
        for job in self.db.rows("SELECT funnel,created_at FROM jobs WHERE funnel IS NOT NULL AND funnel<>'' "
                                "ORDER BY id DESC LIMIT 30"):
            funnel = load_funnel(job['funnel']) or {}
            for entry in (funnel.get('accounts') or {}).values():
                identifier = entry.get('identifier')
                if not identifier or identifier in found:
                    continue
                failed = entry.get('status') == 'failed'
                returned = (entry.get('counts') or {}).get('returned', 0)
                state = '上次采集失败' if failed else f"采集正常 · 上次返回 {returned} 篇"
                if entry.get('latest_published') and not failed:
                    state += '，最近一篇 ' + china_time(entry['latest_published'])
                found[identifier] = {'state': state, 'failed': failed, 'error': entry.get('error') or '',
                                     'coverage': format_coverage(entry.get('coverage'), returned),
                                     'when': china_time(job['created_at'])}
        return found

    def account_detail(self, account, details):
        info = details.get(account['identifier'], {})
        win = tk.Toplevel(self.root)
        win.title(account['name'])
        win.configure(bg=self.pal['paper'], padx=20, pady=16)
        win.transient(self.root)
        lines = [f"名称：{account['name']}", f"标识：{account['identifier']}",
                 f"状态：{'启用' if account['enabled'] else '停用'}"]
        if info:
            lines.append(f"最近一次采集（{info['when']}）：{info['state']}")
            lines.append('这是本软件那次采集实际返回的结果，不代表该公众号在全网的最近发文。')
            if info.get('coverage'):
                lines.append(info['coverage'])
            if info.get('error'):
                lines.append('失败原因：' + redact(info['error']))
        else:
            lines.append('还没有采集记录。')
        for line in lines:
            ttk.Label(win, text=line, style='Card.TLabel', wraplength=560).pack(anchor='w', pady=2)
        bar = tk.Frame(win, bg=self.pal['paper'])
        bar.pack(fill='x', pady=(14, 0))

        def delete():
            self.require_idle()
            if not messagebox.askyesno('删除公众号', f"将 {account['name']} 移出后续采集名单。\n"
                                       '已采集的文章和已生成的稿件都会保留。是否删除？', parent=win):
                return
            self.db.execute("DELETE FROM accounts WHERE id=?", (account['id'],))
            win.destroy()
            self.show('公众号')
        ttk.Button(bar, text='删除', style='Danger.TButton', command=lambda: self.safe(delete)).pack(side='left')
        ttk.Button(bar, text='关闭', command=win.destroy).pack(side='right')

    def weread_source(self, cancel=None):
        from .weread import WeReadSource
        return WeReadSource(self.vault.get('weread'), self.db, cancel,
                            state_saver=lambda value: self.vault.set_many({'weread': value}),
                            profile_dir=self.data_dir / 'weread-edge-profile',
                            skill_key=(self.vault.get('weread_skill')
                                       if self.vault.has('weread_skill') else None))

    def open_sync(self):
        """扫码同步书架 opens an inline panel: login if needed → read shelf → pick → import."""
        self.require_live()
        self.require_idle()
        self.sync_state = {'phase': 'start', 'cancel': threading.Event()}
        self.render_sync()
        self.begin_sync(force_login=False)

    def close_sync(self):
        state = getattr(self, 'sync_state', None)
        if state:
            state['cancel'].set()
        self.sync_state = None
        if self.page == '公众号':
            for child in self.sync_holder.winfo_children():
                child.destroy()

    def begin_sync(self, force_login):
        from .weread import browser_login
        state = self.sync_state
        need_login = force_login or not self.vault.has('weread')
        state.update(phase='login' if need_login else 'reading', message='')
        self.render_sync()
        cancel = state['cancel']

        def work():
            if need_login:
                login_state, _ = browser_login(self.data_dir / 'weread-edge-profile', cancel)
                self.vault.set_many({'weread': login_state})
            source = self.weread_source(cancel)
            try:
                return source.shelf()
            finally:
                source.close()

        def done(books):
            if self.sync_state is state:
                state.update(phase='choose', books=books)
                self.render_sync()

        def failed(exc):
            if self.sync_state is not state:
                return
            auth = isinstance(exc, ApiError) and exc.kind == 'weread_auth'
            state.update(phase='error', relogin=auth or need_login,
                         message=('登录已失效：' if auth else '读取书架失败：') + redact(str(exc)) + '。现有公众号名单不变。')
            self.render_sync()
        self.start_worker(work, '正在同步微信读书书架…', on_success=done, on_error=failed)

    def render_sync(self):
        from .weread import shelf_accounts, sync_books
        holder, state = self.sync_holder, self.sync_state
        for child in holder.winfo_children():
            child.destroy()
        card = self.card(holder, title='同步微信读书书架')
        self.pack_button(card.head, '关闭', self.close_sync, side='right')
        self.text(card, '仅同步已加入微信读书书架的公众号（不是微信关注列表）。', 'muted', anchor='w', pady=(0, 10))
        phase = state['phase']
        if phase in ('start', 'login', 'reading'):
            box = tk.Frame(card, bg=self.pal['base'], width=260, height=110)
            box.pack(anchor='w')
            box.pack_propagate(False)
            text = ('已打开 Edge 登录窗口\n请用手机微信扫码确认' if phase == 'login' else '正在读取微信读书书架…')
            tk.Label(box, text=text, bg=self.pal['base'], fg=self.pal['muted'], font=(FONT, 9),
                     justify='center').pack(expand=True)
            self.text(card, '二维码在弹出的 Edge 窗口中；扫码成功后会自动读取书架。二维码过期时关闭此面板后重新点击即可。',
                      'muted', 9, anchor='w', pady=(8, 0), wrap=700)
            return
        if phase == 'error':
            self.text(card, state['message'], 'warn', anchor='w', wrap=760)
            line = self.row(card, pady=(10, 0))
            if state.get('relogin'):
                self.pack_button(line, '重新扫码登录', lambda: self.begin_sync(True), 'primary')
            self.pack_button(line, '重试读取', lambda: self.begin_sync(False))
            return
        books = state['books']
        shelf = shelf_accounts(books)
        existing = {a['identifier'] for a in self.db.accounts()}
        if not shelf:
            self.text(card, '书架里没有公众号。请先在手机微信读书把公众号加入书架，再回来同步；现有名单不变。', 'muted',
                      anchor='w')
            return
        fresh = [a for a in shelf if a['identifier'] not in existing]
        self.text(card, f"书架共 {len(shelf)} 个公众号，新增 {len(fresh)} 个。", 'ink', anchor='w', pady=(0, 6))
        marks = state.setdefault('marks', {})
        for account in shelf:
            added = account['identifier'] in existing
            var = marks.setdefault(account['identifier'], tk.BooleanVar(value=not added))
            check = ttk.Checkbutton(card, text=account['name'] + ('（已在名单中）' if added else ''), variable=var)
            check.pack(anchor='w')
            if added:
                check.state(['disabled'])
        line = self.row(card, pady=(10, 0))

        def select_all():
            for identifier, var in marks.items():
                if identifier not in existing:
                    var.set(True)

        def run_import():
            chosen = {identifier for identifier, var in marks.items() if var.get() and identifier not in existing}
            if not chosen:
                raise ValueError('请至少勾选一个新增公众号')
            result = sync_books(self.db, books, selected=chosen)
            self.sync_state = None
            self.status.set(result)
            self.show('公众号')
        self.pack_button(line, '全选新增公众号', select_all)
        self.pack_button(line, '导入所选', run_import, 'primary')

    def add_account_dialog(self):
        from .weread import identifier_from_link
        win = tk.Toplevel(self.root)
        win.title('添加公众号')
        win.configure(bg=self.pal['paper'], padx=20, pady=16)
        win.transient(self.root)
        name, link = tk.StringVar(), tk.StringVar()
        ttk.Label(win, text='公众号名称', style='Card.TLabel').pack(anchor='w')
        name_entry = ttk.Entry(win, textvariable=name, width=46)
        name_entry.pack(fill='x', pady=(2, 10))
        ttk.Label(win, text='该公众号任一文章链接，或 MP_WXS_ 开头的标识', style='Card.TLabel').pack(anchor='w')
        ttk.Entry(win, textvariable=link, width=46).pack(fill='x', pady=(2, 6))
        ttk.Label(win, text='软件从链接里的 __biz 识别账号，不按名称猜测；名称需与公众号实际名称一致，用于向采集服务查询。',
                  style='Muted.TLabel', wraplength=420).pack(anchor='w', pady=(0, 10))
        bar = tk.Frame(win, bg=self.pal['paper'])
        bar.pack(fill='x')

        def add():
            self.require_idle()
            identifier = identifier_from_link(link.get())
            if self.db.rows('SELECT id FROM accounts WHERE identifier=?', (identifier,)):
                raise ValueError('该公众号已在名单中')
            self.db.add_account(name.get(), identifier)
            win.destroy()
            self.status.set(f"已添加 {name.get().strip()}，从下一个新任务起采集")
            self.show('公众号')
        ttk.Button(bar, text='添加', style='Primary.TButton', command=lambda: self.safe(add)).pack(side='right')
        ttk.Button(bar, text='取消', command=win.destroy).pack(side='right', padx=8)
        name_entry.focus_set()

    # --------------------------------------------------------------- articles

    def articles(self):
        self.heading('文章库')
        bar = ttk.Frame(self.content)
        bar.pack(fill='x', pady=(0, 8))
        query = tk.StringVar()
        entry = ttk.Entry(bar, textvariable=query)
        names = ['全部公众号'] + [r['account_name'] for r in self.db.rows(
            "SELECT DISTINCT account_name FROM articles WHERE cleared=0 ORDER BY account_name")]
        account, period, order = tk.StringVar(value='全部公众号'), tk.StringVar(value='全部时间'), tk.StringVar(value='阅读量优先')
        for var, values, width in ((order, ['阅读量优先', '点赞量优先', '最新发布'], 10),
                                   (period, ['全部时间', '今天', '近 3 天', '近 7 天'], 9),
                                   (account, names, 14)):
            self.combo(bar, var, values, width).pack(side='right', padx=(8, 0))
        entry.pack(side='left', fill='x', expand=True)
        self.placeholder(entry, query, '搜索文章标题')
        meta = ttk.Frame(self.content)
        meta.pack(fill='x', pady=(0, 8))
        count_text = tk.StringVar()
        ttk.Label(meta, textvariable=count_text, style='Base.Muted.TLabel').pack(side='left')
        ttk.Label(meta, text='发布时间 · 中国时间', style='Base.Muted.TLabel').pack(side='right')
        holder = tk.Frame(self.content, bg=self.pal['base'])
        holder.pack(fill='x')
        detail = tk.Frame(self.content, bg=self.pal['base'])
        detail.pack(fill='x')
        state = {'page': 0}
        used = self.engine._used_articles()
        size = 20

        def render(line, a):
            left = tk.Frame(line, bg=self.pal['paper'])
            left.pack(side='left', fill='x', expand=True)
            self.text(left, a['title'], 'ink', 10, wrap=760)
            sub = tk.Frame(left, bg=self.pal['paper'])
            sub.pack(anchor='w', pady=(3, 2))
            self.text(sub, a['account_name'], 'muted', 9, side='left', padx=(0, 10))
            self.text(sub, china_time(a['published_at']), 'muted', 9, side='left', padx=(0, 10))
            status, color = (('已用于生成', 'accent') if a['id'] in used else ('正文失败', 'bad') if a['error'] else
                             ('正文已获取', 'muted') if a['has_body'] else ('正文未获取', 'warn'))
            if a['id'] not in used and not a['error'] and a['has_body']:
                status = '未使用'
            self.text(sub, status, color, 9, side='left')
            stats = tk.Frame(left, bg=self.pal['paper'])
            stats.pack(anchor='w')
            for label, value in (('阅读', a['reads']), ('点赞', a['likes'])):
                self.text(stats, label, 'muted', 9, side='left', padx=(0, 4))
                self.text(stats, '未知' if value is None else f'{value:,}', 'ink', 9, side='left', padx=(0, 14))
            self.pack_button(line, '查看文章' if a['has_body'] else '查看详情',
                             lambda i=a['id']: open_detail(i), side='right')

        def open_detail(article_id):
            self.article_open = article_id
            draw_detail()

        def draw_detail():
            for child in detail.winfo_children():
                child.destroy()
            article_id = getattr(self, 'article_open', None)
            if not article_id:
                return
            found = self.db.rows("SELECT * FROM articles WHERE id=?", (article_id,))
            if not found:
                return
            article = found[0]
            card = self.card(detail, title=article['title'])
            self.pack_button(card.head, '收起', lambda: (setattr(self, 'article_open', None), draw_detail()),
                             side='right')
            self.text(card, f"{article['account_name']} · {china_time(article['published_at'], '%Y-%m-%d %H:%M')}（中国时间）",
                      'muted', 9, anchor='w', pady=(0, 10))
            if article['error']:
                self.text(card, '正文处理失败：' + redact(article['error']), 'bad', anchor='w', wrap=860, pady=(0, 8))
            body = ScrolledText(card, wrap='word', height=12, font=(FONT, 10), relief='flat', padx=4, pady=4,
                                bg=self.pal['paper'], fg=self.pal['ink'])
            body.pack(fill='x')
            images = self.db.article_images(article['id'])
            self.fill_text(body, article['content'] or '尚未获取正文（任务选中这篇文章时才会读取）。', images)
            body.configure(state='disabled')
            line = self.row(card, pady=(10, 0))
            self.pack_button(line, '打开原文', lambda: self.open_article(article_id))
            self.rule(card, pady=(12, 0))

            def collect_info(box):
                metrics = self.db.rows("SELECT * FROM article_metrics WHERE article_id=?", (article_id,))
                note = self.db.rows("SELECT source FROM metric_notes WHERE article_id=? ORDER BY rowid DESC LIMIT 1",
                                    (article_id,))
                lines = [f"原文：{article['url']}",
                         f"采集时间：{china_time(article['fetched_at'], '%Y-%m-%d %H:%M')}"]
                if metrics:
                    lines.append(f"指标采集时间：{china_time(metrics[0]['observed_at'], '%Y-%m-%d %H:%M')}")
                if note and note[0]['source']:
                    lines.append('指标来源：' + note[0]['source'])
                lines.append(f'原文图片：已保存 {len(images)} 张')
                for text in lines:
                    self.text(box, text, 'muted', 9, anchor='w', wrap=860)
                analyses = self.db.rows("SELECT result FROM analyses WHERE article_id=? ORDER BY created_at",
                                        (article_id,))
                if analyses:
                    ttk.Button(box, text='查看摘要与证据', command=lambda: self.text_window(
                        '摘要与证据', '\n\n'.join(json.dumps(json.loads(a['result']), ensure_ascii=False, indent=2)
                                              for a in analyses))).pack(anchor='w', pady=(6, 0))
            self.toggle(card, '采集详情', on_open=collect_info)
            self.root.after(50, lambda: self.canvas.yview_moveto(1.0))

        def refresh(*_):
            where, params = ["a.cleared=0"], []
            text = query.get().strip()
            if text and text != '搜索文章标题':
                pattern = '%' + text.replace('!', '!!').replace('%', '!%').replace('_', '!_') + '%'
                where.append("a.title LIKE ? ESCAPE '!'")
                params.append(pattern)
            if account.get() != '全部公众号':
                where.append('a.account_name=?')
                params.append(account.get())
            now = datetime.now(CHINA)
            since = {'今天': now.replace(hour=0, minute=0, second=0, microsecond=0),
                     '近 3 天': now - timedelta(days=3), '近 7 天': now - timedelta(days=7)}.get(period.get())
            if since:
                where.append('a.published_at>=?')
                params.append(since.astimezone(timezone.utc).isoformat())
            # Unknown metrics sort after every valid value (0 included); the
            # whole result is ordered before paging.
            sort = {'阅读量优先': '(m.reads IS NULL), m.reads DESC, a.published_at DESC',
                    '点赞量优先': '(m.likes IS NULL), m.likes DESC, a.published_at DESC',
                    '最新发布': 'a.published_at DESC'}[order.get()]
            sql_where = ' WHERE ' + ' AND '.join(where)
            total = self.db.rows('SELECT COUNT(*) n FROM articles a' + sql_where, params)[0]['n']
            pages = max(1, (total + size - 1) // size)
            state['page'] = max(0, min(state['page'], pages - 1))
            rows = self.db.rows(
                "SELECT a.id,a.title,a.account_name,a.published_at,a.content<>'' has_body,a.error,m.reads,m.likes "
                "FROM articles a LEFT JOIN article_metrics m ON m.article_id=a.id" + sql_where +
                " ORDER BY " + sort + " LIMIT ? OFFSET ?", params + [size, state['page'] * size])
            count_text.set(f'共 {total} 篇')
            for child in holder.winfo_children():
                child.destroy()
            self.list_rows(holder, rows, render, '没有符合条件的文章。')
            self.pager(holder, state, total, size, refresh)
        for var in (query, account, period, order):
            var.trace_add('write', lambda *_: (state.update(page=0), refresh()))
        refresh()
        draw_detail()

    def open_article(self, article_id):
        url = self.db.rows("SELECT url FROM articles WHERE id=?", (article_id,))[0]["url"]
        self.open_link(wechat_url(url))

    def article_detail(self, article_id):
        """Open an article's inline detail on the 文章库 page."""
        if self.page != '文章库':
            self.show('文章库')
        self.article_open = article_id
        self.show('文章库')

    # ----------------------------------------------------------------- drafts

    def open_draft(self, draft_id):
        if draft_id is None:
            return
        self.show('稿件库')
        self.draft_id = draft_id
        self.show_draft()

    def drafts(self):
        if self.draft_id:
            self.show_draft()
            return
        self.heading('稿件库')
        bar = ttk.Frame(self.content)
        bar.pack(fill='x', pady=(0, 8))
        query = tk.StringVar(value=getattr(self, 'draft_query', ''))
        status = tk.StringVar(value=getattr(self, 'draft_filter', '全部状态'))
        self.combo(bar, status, ['全部状态', '待审核', '需修改', '已通过'], 10).pack(side='right', padx=(8, 0))
        entry = ttk.Entry(bar, textvariable=query)
        entry.pack(side='left', fill='x', expand=True)
        self.placeholder(entry, query, '搜索稿件标题')
        count_text = tk.StringVar()
        ttk.Label(self.content, textvariable=count_text, style='Base.Muted.TLabel').pack(anchor='w', pady=(0, 8))
        holder = tk.Frame(self.content, bg=self.pal['base'])
        holder.pack(fill='x')
        state = {'page': getattr(self, 'draft_page', 0)}
        size = 20

        def render(line, draft):
            review = Database.review_state(draft['status'])
            left = tk.Frame(line, bg=self.pal['paper'])
            left.pack(side='left', fill='x', expand=True)
            self.text(left, draft['title'], 'ink', 10, wrap=760)
            sub = tk.Frame(left, bg=self.pal['paper'])
            sub.pack(anchor='w', pady=(3, 0))
            self.text(sub, review, {'已通过': 'good', '需修改': 'warn'}.get(review, 'muted'), 9, side='left',
                      padx=(0, 12))
            self.text(sub, china_time(draft['updated_at']) + '（中国时间）', 'muted', 9, side='left', padx=(0, 12))
            hint = draft_hint(draft['report'])
            if hint:
                self.text(sub, hint, 'warn', 9, side='left')
            self.pack_button(line, '查看稿件', lambda d=draft['id']: self.open_draft(d), side='right')

        def refresh(*_):
            text = query.get().strip()
            self.draft_query = '' if text == '搜索稿件标题' else text
            self.draft_filter = status.get()
            where, params = ["cleared=0"], []
            if self.draft_query:
                pattern = '%' + self.draft_query.replace('!', '!!').replace('%', '!%').replace('_', '!_') + '%'
                where.append("(title LIKE ? ESCAPE '!' OR body LIKE ? ESCAPE '!')")
                params += [pattern, pattern]
            if self.draft_filter != '全部状态':
                states = ('已通过', '已审核') if self.draft_filter == '已通过' else (self.draft_filter,)
                where.append('status IN (' + ','.join('?' * len(states)) + ')')
                params += list(states)
            sql_where = ' WHERE ' + ' AND '.join(where)
            total = self.db.rows('SELECT COUNT(*) n FROM drafts' + sql_where, params)[0]['n']
            pages = max(1, (total + size - 1) // size)
            state['page'] = max(0, min(state['page'], pages - 1))
            self.draft_page = state['page']
            rows = self.db.rows('SELECT id,title,status,updated_at,report FROM drafts' + sql_where +
                                ' ORDER BY id DESC LIMIT ? OFFSET ?', params + [size, state['page'] * size])
            count_text.set(f'共 {total} 篇 · 最新生成优先')
            for child in holder.winfo_children():
                child.destroy()
            self.list_rows(holder, rows, render, '没有符合条件的稿件。')
            self.pager(holder, state, total, size, refresh)
        query.trace_add('write', lambda *_: (state.update(page=0), refresh()))
        status.trace_add('write', lambda *_: (state.update(page=0), refresh()))
        refresh()

    def show_draft(self):
        found = self.db.rows('SELECT * FROM drafts WHERE id=? AND cleared=0', (self.draft_id,))
        for child in self.content.winfo_children():
            child.destroy()
        if not found:
            self.draft_id = None
            self.drafts()
            return
        draft = found[0]
        review = Database.review_state(draft['status'])
        p = self.pal

        def back():
            if self.editor_dirty and not messagebox.askyesno('未保存修改', '返回后将丢弃未保存的修改，是否返回？',
                                                             parent=self.root):
                return
            self.editor_dirty = False
            self.draft_id = None
            self.show('稿件库')
        bar = ttk.Frame(self.content)
        bar.pack(fill='x', pady=(0, 12))
        ttk.Button(bar, text='← 返回稿件库', command=lambda: self.safe(back)).pack(side='left')
        tk.Label(bar, text=review + ('（人工已确认，未发布）' if review == '已通过' else ''), bg=p['base'],
                 fg={'已通过': p['good'], '需修改': p['warn']}.get(review, p['muted']), font=(FONT, 9)).pack(side='right')
        card = self.card()
        title = tk.StringVar(value=draft['title'])
        title_entry = tk.Entry(card, textvariable=title, font=(FONT, 15, 'bold'), relief='flat', bd=0,
                               bg=p['paper'], fg=p['ink'], insertbackground=p['ink'], readonlybackground=p['paper'],
                               disabledforeground=p['ink'])
        title_entry.pack(fill='x')
        self.text(card, f"{china_time(draft['updated_at'])}（中国时间） · 版本 {draft['revision']}", 'muted', 9,
                  anchor='w', pady=(4, 8))
        hint = draft_hint(draft['report'])
        if hint:
            self.text(card, hint + '，发布前请展开“来源与审核提示”核对。', 'warn', anchor='w', pady=(0, 8))
        editor = ScrolledText(card, wrap='word', undo=True, font=(FONT, 11), relief='flat',
                              height=14 if self.compact else 20, padx=0, pady=6, bd=0,
                              bg=p['paper'], fg=p['ink'], insertbackground=p['ink'])
        editor.pack(fill='both', expand=True)
        editor.insert('1.0', draft['body'])
        editor.edit_modified(False)
        editing = {'on': False}
        actions = self.row(card, pady=(12, 0))
        note = tk.StringVar(value='')
        note_label = ttk.Label(card, textvariable=note, style='Muted.TLabel')
        note_label.pack(anchor='w', pady=(6, 0))

        def lock(on):
            editing['on'] = on
            editor.configure(state='normal' if on else 'disabled',
                             bg=p['field'] if on else p['paper'], padx=10 if on else 0)
            title_entry.configure(state='normal' if on else 'readonly',
                                  bg=p['field'] if on else p['paper'])
            for button in all_buttons:
                button.pack_forget()
            shown = [save_button, cancel_button] if on else [edit_button, *review_buttons, export_button]
            for button in shown:
                button.pack(side='left', padx=(0, 8), pady=2)
            note.set('保存修改后需重新审核。' if on else '')
        editor.bind('<<Modified>>', lambda e: (setattr(self, 'editor_dirty', editing['on'] or self.editor_dirty),
                                                editor.edit_modified(False)))
        title.trace_add('write', lambda *_: setattr(self, 'editor_dirty', True) if editing['on'] else None)

        def save():
            self.require_idle()
            body = editor.get('1.0', 'end-1c')
            self.db.edit_draft(draft['id'], title.get(), body, draft['status'])
            self.editor_dirty = False
            changed = (title.get().strip(), body.strip()) != (draft['title'], draft['body'])
            self.status.set('已保存修改' + ('，稿件回到待审核' if changed and review != '待审核' else ''))
            self.show_draft()

        def cancel_edit():
            if self.editor_dirty and not messagebox.askyesno('取消编辑', '放弃本次修改？', parent=self.root):
                return
            self.editor_dirty = False
            self.show_draft()

        def mark(state_name, ask=None):
            self.require_idle()
            if ask and not messagebox.askyesno('确认', ask, parent=self.root):
                return
            self.db.set_draft_status(draft['id'], state_name)
            self.status.set(f'已标记为“{state_name}”，审核记录已写入素材报告')
            self.show_draft()
        edit_button = self.button(actions, '编辑', lambda: lock(True))
        save_button = self.button(actions, '保存修改', save, 'primary')
        cancel_button = self.button(actions, '取消编辑', cancel_edit)
        review_buttons = [self.button(actions, '标记需修改', lambda: mark('需修改')),
                          self.button(actions, '审核通过', lambda: mark(
                              '已通过', '已核对正文、事实与来源，确认审核通过？\n审核通过不会自动发布。'), 'primary')]
        export_box = tk.Frame(card, bg=p['paper'])
        export_button = self.button(actions, '导出', lambda: (
            export_box.pack(fill='x', pady=(10, 0), after=actions) if not export_box.winfo_ismapped()
            else export_box.pack_forget()))
        all_buttons = [edit_button, save_button, cancel_button, *review_buttons, export_button]
        lock(False)

        config = self.db.config()
        fmt = tk.StringVar(value=EXPORT_FORMATS[config.get('export_format', 'docx')])
        line = tk.Frame(export_box, bg=p['paper'])
        line.pack(fill='x')
        self.text(line, '格式', 'muted', 9, side='left', padx=(0, 8))
        self.combo(line, fmt, list(EXPORT_FORMATS.values()), 16).pack(side='left', padx=(0, 12))
        code = lambda: next(k for k, v in EXPORT_FORMATS.items() if v == fmt.get())

        def target(default_name):
            if self.editor_dirty:
                raise ValueError('请先保存或取消当前修改，再导出')
            suffix = '.' + code()
            folder = config.get('export_dir') or ''
            initial = re.sub(r'[\\/:*?"<>|]+', '-', default_name)[:80] + suffix
            if folder and Path(folder).is_dir():
                path = Path(folder) / initial
                if path.exists() and not messagebox.askyesno('覆盖文件', f'{path.name} 已存在，是否覆盖？', parent=self.root):
                    return None
                return path
            chosen = filedialog.asksaveasfilename(parent=self.root, defaultextension=suffix, initialfile=initial,
                                                  filetypes=[(fmt.get(), '*' + suffix)])
            return Path(chosen) if chosen else None

        def export_copy():
            current = self.db.rows('SELECT * FROM drafts WHERE id=?', (draft['id'],))[0]
            if Database.review_state(current['status']) != '已通过' and not messagebox.askyesno(
                    '稿件未审核通过', '这篇稿件还没有人工审核通过，仍要导出吗？', parent=self.root):
                return
            path = target(current['title'])
            if path is None:
                self.status.set('已取消导出')
                return
            write_export(path, code(), current['title'], current['body'], json.loads(current['sources']))
            self.status.set(f'已导出发布稿：{path}')

        def export_report():
            current = self.db.rows('SELECT * FROM drafts WHERE id=?', (draft['id'],))[0]
            path = target('素材报告-' + current['title'])
            if path is None:
                self.status.set('已取消导出')
                return
            write_report(path, code(), current['title'], current['report'])
            self.status.set(f'已导出素材报告：{path}')
        self.pack_button(line, '导出发布稿', export_copy, 'primary')
        self.pack_button(line, '导出素材报告', export_report)
        self.text(export_box, '发布稿不含来源编号、审核备注；来源对应关系保留在素材报告中。', 'muted', 9, anchor='w',
                  pady=(6, 0))
        self.rule(card, pady=(14, 0))

        def report_body(body):
            text = ScrolledText(body, wrap='word', height=12, font=(FONT, 10), relief='flat', padx=10, pady=8,
                                bg=p['base'], fg=p['ink'])
            text.pack(fill='both', expand=True)
            text.insert('1.0', redact(draft['report']))
            text.configure(state='disabled')
            tools = tk.Frame(body, bg=p['paper'])
            tools.pack(fill='x', pady=(6, 0))
            ttk.Button(tools, text='参考网站', command=lambda: self.safe(lambda: self.text_window(
                '参考资料网站', source_website_text(draft['body'], json.loads(draft['sources']))))).pack(side='left')
            ttk.Button(tools, text='历史版本', command=lambda: self.safe(versions)).pack(side='left', padx=8)

        def versions():
            rows = self.db.rows("SELECT * FROM draft_versions WHERE draft_id=? ORDER BY revision DESC", (draft['id'],))
            self.text_window("稿件历史版本", "\n\n".join(
                f"版本 {v['revision']} · {china_time(v['created_at'])}\n{v['title']}\n\n{v['body']}" for v in rows))
        self.toggle(card, '来源与审核提示', on_open=report_body)

    # --------------------------------------------------------------- settings

    def settings(self):
        config = self.db.config()
        p = self.pal
        variables, keys, original = {}, {}, {}

        def track(name, var):
            variables[name] = var
            original[name] = var.get()
            var.trace_add('write', lambda *_: mark_dirty())
            return var

        ready = {'on': False}

        def mark_dirty():
            if not ready['on']:
                return
            self.settings_dirty = (any(var.get() != original.get(name) for name, var in variables.items())
                                   or any(var.get().strip() for var in keys.values())
                                   or extra_text.get('1.0', 'end-1c').strip() != original.get('style_extra', ''))
            for button in (save_button, undo_button):
                if button.winfo_exists():
                    button.state(['!disabled'] if self.settings_dirty else ['disabled'])
            saved_note.set('' if self.settings_dirty else '设置已保存')
        self.heading('设置')

        def key_entry(parent, name, hint):
            line = tk.Frame(parent, bg=p['paper'])
            var = tk.StringVar()
            entry = ttk.Entry(line, textvariable=var, show='●')
            entry.pack(side='left', fill='x', expand=True)
            keys[name] = var
            var.trace_add('write', lambda *_: mark_dirty())
            shown = {'on': False}

            def reveal():
                shown['on'] = not shown['on']
                entry.configure(show='' if shown['on'] else '●')
                toggle_button.configure(text='隐藏' if shown['on'] else '显示')
            toggle_button = ttk.Button(line, text='显示', command=reveal)
            toggle_button.pack(side='left', padx=(8, 0))
            saved = self.vault.has(name)
            tk.Label(parent, text=('已配置；留空保留原 Key' if saved else '未配置') + (' · ' + hint if hint else ''),
                     bg=p['paper'], fg=p['good'] if saved else p['warn'], font=(FONT, 8)).pack(anchor='w', side='bottom')
            return line

        # Model: one writing model in front (provider follows the model ID).
        models = self.card(title='模型')
        model_vars = {name: track(name + '_model', tk.StringVar(value=config['providers'][name]['model']))
                      for name in ('openai', 'deepseek')}
        writer = track('writer_provider', tk.StringVar(value=config['writer_provider']))
        choice = tk.StringVar(value=model_vars[writer.get()].get())

        def picked(*_):
            model = choice.get().strip()
            if not model:
                return
            provider = 'deepseek' if model.lower().startswith('deepseek') else 'openai'
            writer.set(provider)
            model_vars[provider].set(model)
            key_title.configure(text=('DeepSeek' if provider == 'deepseek' else 'OpenAI') + ' API Key')
            for name, frame in key_frames.items():
                frame.pack_forget()
            key_frames[provider].pack(fill='x')
        options = MODEL_OPTIONS['deepseek'] + MODEL_OPTIONS['openai']
        if choice.get() not in options:
            options = [choice.get()] + options
        self.field(models, '写作与二审', lambda box: ttk.Combobox(box, textvariable=choice, values=options, width=40))
        key_title = tk.Label(models, text='', bg=p['paper'], fg=p['ink'], font=(FONT, 9, 'bold'))
        key_title.pack(anchor='w', pady=(0, 4))
        key_box = tk.Frame(models, bg=p['paper'])
        key_box.pack(fill='x', pady=(0, 6))
        key_frames = {}
        for name, hint in (('deepseek', ''), ('openai', '联网搜索也使用这个 Key')):
            holder = tk.Frame(key_box, bg=p['paper'])
            key_entry(holder, name, hint).pack(fill='x')
            key_frames[name] = holder
        choice.trace_add('write', picked)
        picked()
        for name in ('summary_provider', 'search_model'):
            track(name, tk.StringVar(value=config[name]))
        for name, value in (('openai_url', config['providers']['openai']['base_url']),
                            ('deepseek_url', config['providers']['deepseek']['base_url']),
                            ('search_url', config['search_url'])):
            track(name, tk.StringVar(value=value))

        def addresses(body):
            for label, name in (('OpenAI 接口地址', 'openai_url'), ('DeepSeek 接口地址', 'deepseek_url'),
                                ('联网研究接口', 'search_url')):
                self.field(body, label, lambda box, n=name: ttk.Entry(box, textvariable=variables[n]))

        def more_models(body):
            self.field(body, '摘要与分析使用', lambda box: self.combo(box, variables['summary_provider'],
                                                                ['deepseek', 'openai'], 12))
            self.field(body, '联网搜索模型（OpenAI web_search）', lambda box: ttk.Combobox(
                box, textvariable=variables['search_model'], values=MODEL_OPTIONS['openai'], width=30))
            other = 'openai' if writer.get() == 'deepseek' else 'deepseek'
            self.field(body, ('OpenAI' if other == 'openai' else 'DeepSeek') + ' 模型（摘要或备用）',
                       lambda box: ttk.Combobox(box, textvariable=model_vars[other], values=MODEL_OPTIONS[other],
                                                width=30))
            label = ('OpenAI' if other == 'openai' else 'DeepSeek') + ' API Key'
            saved = '已配置；留空保留原 Key' if self.vault.has(other) else '未配置'
            self.field(body, f'{label}（{saved}）', lambda box: ttk.Entry(box, textvariable=keys[other], show='●'))

            def recommended():
                choice.set(RECOMMENDED['deepseek'])
                variables['summary_provider'].set(RECOMMENDED['summary_provider'])
                variables['search_model'].set(RECOMMENDED['search_model'])
                self.status.set('已填入推荐模型（DeepSeek deepseek-v4-pro 写作与摘要，gpt-5.6-terra 联网搜索），点“保存设置”后生效')
            ttk.Button(body, text='填入推荐模型', command=lambda: self.safe(recommended)).pack(anchor='w')
        self.toggle(models, '接口地址', on_open=addresses)
        self.toggle(models, '更多模型配置', on_open=more_models)
        self.text(models, '模型选项以实际接入服务为准，可直接输入服务支持的其他模型 ID。推荐：文档处理 deepseek-v4-pro，'
                          '联网搜索 gpt-5.6-terra（更省钱可选 gpt-6-luna）。', 'muted', 9, anchor='w',
                  pady=(8, 0), wrap=860)

        # Writing preferences: two columns, then the free-text requirement.
        writing = self.card(title='写作偏好')
        grid = tk.Frame(writing, bg=p['paper'])
        grid.pack(fill='x')
        grid.columnconfigure(0, weight=1, uniform='half')
        grid.columnconfigure(1, weight=1, uniform='half')
        words = track('target_words', tk.StringVar(value=str(config['target_words'])))
        self.field(grid, '目标篇幅（字，可填 200—20000 的任意整数）', lambda box: ttk.Combobox(
            box, textvariable=words, values=['800', '1200', '1500', '2000', '3000']), column=0)
        preset_value = config.get('style_preset') or next(
            (name for name, text in STYLE_PRESETS.items() if text == config['style']), config['style'])
        preset = track('style_preset', tk.StringVar(value=preset_value))
        self.field(grid, '写作风格', lambda box: ttk.Combobox(box, textvariable=preset, values=list(STYLE_PRESETS)),
                   column=1)
        original['style_extra'] = config.get('style_extra', '') if config.get('style_preset') else ''
        self.field(writing, '补充要求', lambda box: tk.Text(box, height=4, wrap='word', relief='solid', bd=1,
                                                         font=(FONT, 10), bg=p['field'], fg=p['ink'],
                                                         insertbackground=p['ink'], highlightthickness=0))
        extra_text = [w for w in writing.winfo_children()[-1].winfo_children() if isinstance(w, tk.Text)][0]
        extra_text.insert('1.0', original['style_extra'])
        extra_text.bind('<KeyRelease>', lambda e: mark_dirty())

        output = self.card(title='导出')
        grid = tk.Frame(output, bg=p['paper'])
        grid.pack(fill='x')
        grid.columnconfigure(0, weight=1, uniform='half')
        grid.columnconfigure(1, weight=1, uniform='half')
        fmt = track('export_format', tk.StringVar(value=EXPORT_FORMATS[config.get('export_format', 'docx')]))
        self.field(grid, '默认格式', lambda box: self.combo(box, fmt, list(EXPORT_FORMATS.values()), 20), column=0)
        folder = track('export_dir', tk.StringVar(value=config.get('export_dir', '')))

        def folder_row(box):
            line = tk.Frame(box, bg=p['paper'])
            shown = tk.Label(line, text=folder.get() or '每次导出时选择', bg=p['paper'], fg=p['muted'], font=(FONT, 9))
            shown.pack(side='left')
            folder.trace_add('write', lambda *_: shown.configure(text=folder.get() or '每次导出时选择'))

            def choose():
                chosen = filedialog.askdirectory(parent=self.root, mustexist=True)
                if chosen:
                    folder.set(chosen)
            ttk.Button(line, text='更改', command=lambda: self.safe(choose)).pack(side='left', padx=(10, 0))
            ttk.Button(line, text='每次选择', style='Link.TButton', command=lambda: folder.set('')).pack(side='left', padx=4)
            return line
        self.field(grid, '保存位置', folder_row, column=1)

        save_row = ttk.Frame(self.content)
        save_row.pack(fill='x', pady=(0, 18))
        save_button = self.button(save_row, '保存设置', lambda: save(), 'primary')
        save_button.pack(side='left')
        undo_button = self.button(save_row, '撤销修改', lambda: undo())
        undo_button.pack(side='left', padx=8)
        saved_note = tk.StringVar(value='设置已保存')
        ttk.Label(save_row, textvariable=saved_note, style='Base.Muted.TLabel').pack(side='left')

        health = self.card()
        line = tk.Frame(health, bg=p['paper'])
        line.pack(fill='x')
        status_vars = {'collect': tk.StringVar(value='待检查'), 'model': tk.StringVar(value='待检查')}
        for name, label in (('collect', '文章采集'), ('model', '文章生成')):
            tk.Label(line, text=label, bg=p['paper'], fg=p['ink'], font=(FONT, 9, 'bold')).pack(side='left')
            tk.Label(line, textvariable=status_vars[name], bg=p['paper'], fg=p['muted'], font=(FONT, 9),
                     wraplength=300, justify='left').pack(side='left', padx=(6, 22))

        def check():
            self.require_live()
            if self.settings_dirty:
                raise ValueError('请先保存设置，再检查连接')
            latest = self.db.config()
            provider = latest['writer_provider']
            billed = messagebox.askyesnocancel(
                '检查连接', '文章采集检查只读取微信读书书架，不计费。\n\n'
                f'是否同时检查文章生成？会向 {provider} 发送一次很短的测试请求，产生少量费用。\n\n'
                '是：两项都检查　否：只检查采集　取消：不检查', parent=self.root)
            if billed is None:
                return
            settings = latest['providers'][provider]
            status_vars['collect'].set('检查中…')
            if billed:
                status_vars['model'].set('检查中…')

            def work():
                results = {}
                try:
                    if not self.vault.has('weread'):
                        raise ValueError('尚未扫码登录微信读书')
                    source = self.weread_source()
                    try:
                        count = len(source.shelf())
                    finally:
                        source.close()
                    services = [label for name, label in (('dajiala', '极致了'), ('everyinfra', 'EveryInfra'))
                                if self.vault.has(name)]
                    results['collect'] = (f"正常（书架 {count} 个公众号；采集服务 Key：{'、'.join(services) or '未配置'}）")
                except Exception as exc:
                    results['collect'] = '异常：' + redact(str(exc))[:120]
                if billed:
                    try:
                        Model(provider, settings, self.vault.get(provider), Transport()).complete(
                            '只输出 JSON 对象 {"ok":true}', "测试连接，输出 JSON", tokens=128)
                        results['model'] = f"正常（{settings['model']}）"
                    except Exception as exc:
                        results['model'] = '异常：' + redact(str(exc))[:120]
                return results

            def done(results):
                stamp = datetime.now(CHINA).strftime('%H:%M')
                for name, value in results.items():
                    status_vars[name].set(f'{value} · {stamp} 检查')
                if not billed:
                    status_vars['model'].set('未检查（计费，需确认后检查）')
            self.start_worker(work, '正在检查连接…', on_success=done)
        self.pack_button(line, '检查连接', check, side='right')

        def help_body(body):
            for text in ('采集异常：先确认已扫码登录；“书架 0 个公众号”说明书架为空。采集服务 Key 是否有效以实际任务为准。',
                         '生成异常：检查 Key、余额和模型 ID；HTTP 401/402/403 多为 Key 或余额问题。',
                         '检查网络后重试；仍不可用时联系服务提供方。'):
                self.text(body, text, 'muted', 9, anchor='w', wrap=860)
        self.toggle(health, '连接异常时怎么办', on_open=help_body)

        services = self.card()

        def service_keys(body):
            self.text(body, '文章列表与阅读量优先使用极致了，失败时改用 EveryInfra；均按次计费。', 'muted', 9, anchor='w',
                      pady=(0, 8))
            for name, label in (('dajiala', '极致了 Key'), ('everyinfra', 'EveryInfra Key（备用）')):
                tk.Label(body, text=label, bg=p['paper'], fg=p['ink'], font=(FONT, 9, 'bold')).pack(anchor='w')
                box = tk.Frame(body, bg=p['paper'])
                box.pack(fill='x', pady=(4, 10))
                key_entry(box, name, '').pack(fill='x')
        self.toggle(services, '采集服务配置', on_open=service_keys)

        maintenance = self.card()
        for name in ('per_account_limit', 'min_reads', 'schedule_time'):
            track(name, tk.StringVar(value=str(config[name])))
        flags = {
            'schedule_enabled': tk.BooleanVar(value=config['schedule_enabled']),
            'quality_review_enabled': tk.BooleanVar(value=config['quality_review_enabled']),
            'image_vision': tk.BooleanVar(value=config.get('image_articles', 'skip') == 'vision'),
            'allow_recollect_cleared': tk.BooleanVar(value=config['allow_recollect_cleared']),
            'startup': tk.BooleanVar(value=startup_enabled()),
        }
        for name, var in flags.items():
            track(name, var)

        def more_settings(body):
            grid = tk.Frame(body, bg=p['paper'])
            grid.pack(fill='x')
            for column, (label, name) in enumerate((('每个公众号列表上限', 'per_account_limit'),
                                                    ('热点最低阅读量', 'min_reads'),
                                                    ('每日执行时间（中国时间）', 'schedule_time'))):
                grid.columnconfigure(column, weight=1, uniform='third')
                box = tk.Frame(grid, bg=p['paper'])
                box.grid(row=0, column=column, sticky='ew', padx=(0, 16))
                tk.Label(box, text=label, bg=p['paper'], fg=p['ink'], font=(FONT, 9, 'bold')).pack(anchor='w', pady=(0, 4))
                ttk.Entry(box, textvariable=variables[name]).pack(fill='x')
            for text, name in (('启用每日自动采集和生成', 'schedule_enabled'),
                               ('生成后进行一次 AI 二审（多一次写作模型调用）', 'quality_review_enabled'),
                               ('识别图片型文章（文字不足 100 字时调用 OpenAI 识图，额外计费）', 'image_vision'),
                               ('允许重新采集已清理的文章', 'allow_recollect_cleared'),
                               ('登录 Windows 时启动软件（托盘运行）', 'startup')):
                ttk.Checkbutton(body, text=text, variable=flags[name]).pack(anchor='w', pady=(6, 0))
            self.text(body, f"当前版本：{version_text()}　数据目录：{self.data_dir}", 'muted', 9, anchor='w',
                      pady=(10, 0), wrap=860)
            line = tk.Frame(body, bg=p['paper'])
            line.pack(fill='x', pady=(6, 0))
            for text, action in (('备份数据库', self.backup), ('恢复数据库', self.restore),
                                 ('打开数据目录', lambda: os.startfile(self.data_dir))):
                ttk.Button(line, text=text, command=lambda a=action: self.safe(a)).pack(side='left', padx=(0, 8))
            self.text(body, '备份包含文章、稿件和配置，不包含 API Key；恢复前会自动保留当前数据库副本，并暂停定时。', 'muted',
                      9, anchor='w', pady=(6, 0))
            line = tk.Frame(body, bg=p['paper'])
            line.pack(fill='x', pady=(6, 0))
            for name, label in (('openai', 'OpenAI'), ('deepseek', 'DeepSeek'), ('dajiala', '极致了'),
                                ('everyinfra', 'EveryInfra')):
                ttk.Button(line, text='清除 ' + label + ' Key', style='Danger.TButton',
                           command=lambda n=name, l=label: self.safe(lambda: clear_key(n, l))).pack(side='left', padx=(0, 8))
        self.toggle(maintenance, '更多设置与维护', on_open=more_settings)

        def clear_key(name, label):
            self.require_idle()
            if not messagebox.askyesno('清除密钥', f'清除已保存的 {label} Key？相关功能将无法使用，直到重新填写。',
                                       parent=self.root):
                return
            self.vault.delete(name)
            self.status.set(label + ' Key 已清除')
            self.settings_dirty = False
            self.show('设置')

        def undo():
            self.settings_dirty = False
            self.show('设置')

        def save():
            self.require_idle()
            self.require_live()
            latest = self.db.config()     # unshown settings are kept as they are
            try:
                latest['target_words'] = int(words.get().strip().replace('约', '').replace('字', '').strip())
                latest['per_account_limit'] = int(variables['per_account_limit'].get().strip())
                latest['min_reads'] = int(variables['min_reads'].get().strip())
            except ValueError:
                raise ValueError('篇幅、列表上限和最低阅读量必须填写整数') from None
            latest['writer_provider'] = writer.get()
            latest['summary_provider'] = variables['summary_provider'].get()
            latest['search_model'] = variables['search_model'].get().strip()
            latest['search_url'] = variables['search_url'].get().strip()
            for provider in ('openai', 'deepseek'):
                latest['providers'][provider] = {
                    'base_url': variables[provider + '_url'].get().strip().rstrip('/'),
                    'model': model_vars[provider].get().strip()}
            latest['style_preset'] = preset.get().strip()
            latest['style_extra'] = extra_text.get('1.0', 'end-1c').strip()
            latest['style'] = compose_style(latest['style_preset'], latest['style_extra'])
            if not latest['style']:
                raise ValueError('请填写写作风格')
            latest['export_format'] = next(k for k, v in EXPORT_FORMATS.items() if v == fmt.get())
            latest['export_dir'] = folder.get()
            latest['schedule_time'] = variables['schedule_time'].get().strip()
            latest['quality_review_enabled'] = flags['quality_review_enabled'].get()
            latest['image_articles'] = 'vision' if flags['image_vision'].get() else 'skip'
            latest['allow_recollect_cleared'] = flags['allow_recollect_cleared'].get()
            latest['comments_enabled'] = False
            if flags['schedule_enabled'].get() and not latest['schedule_enabled']:
                latest['schedule_start'] = datetime.now(ZoneInfo(latest['timezone'])).date().isoformat()
            latest['schedule_enabled'] = flags['schedule_enabled'].get()
            validate(latest)
            self.vault.set_many({name: var.get() for name, var in keys.items() if var.get().strip()})
            if flags['startup'].get() != startup_enabled():
                set_startup(flags['startup'].get())
            self.db.save_config(latest)
            self.settings_dirty = False
            self.status.set('设置已保存；新任务使用新设置，进行中的任务继续用原设置。')
            self.show('设置')
        ready['on'] = True
        mark_dirty()

    def backup(self):
        self.require_idle()
        path = filedialog.asksaveasfilename(parent=self.root, defaultextension=".sqlite", initialfile="公众号工作台备份.sqlite", filetypes=[("SQLite 备份", "*.sqlite")])
        if path:
            self.db.backup(path)
            self.status.set("数据库已备份，不包含 API Key")

    def restore(self):
        self.require_idle()
        path = filedialog.askopenfilename(parent=self.root, filetypes=[("SQLite 备份", "*.sqlite")])
        if not path:
            return
        if not messagebox.askyesno("恢复数据库", "恢复将替换当前文章、稿件和配置，并暂停定时任务。当前数据会先自动备份。是否继续？", parent=self.root):
            return
        self.db.backup(self.data_dir / ("before-restore-" + datetime.now().strftime("%Y%m%d-%H%M%S") + ".sqlite"))
        self.db.restore(path)
        self.status.set("已恢复数据库，定时任务已暂停，请检查配置")
        self.settings_dirty = False
        self.show("设置")

    def run_manual(self):
        """Kept for scripts: start a task with the saved configuration."""
        self.require_live()
        self.start_task(self.db.config())

    # ---------------------------------------------------------------- worker

    def report_callback_exception(self, kind, value, tb):
        logger.error("界面回调异常", exc_info=(kind, value, tb))
        try:
            self.status.set(f"界面出现异常（{kind.__name__}），已记录到日志，可继续使用")
        except Exception:
            pass

    def start_worker(self, action, status, on_success=None, on_error=None):
        self.require_idle()
        self.busy = True
        self.status.set(status)
        self.progress.pack(side="right")
        self.progress.start(12)

        def work():
            try:
                result = action()
                if on_success:
                    self.events.put(("callback", lambda: on_success(result)))
                self.events.put(("done", str(result or "已完成")))
            except Exception as exc:
                safe = isinstance(exc, (ApiError, ValueError, RuntimeError))
                if not isinstance(exc, (ApiError, ValueError)):
                    logger.exception("后台任务异常：%s", status)
                detail = redact(str(exc)) if safe else f"内部错误（{type(exc).__name__}），详情已写入日志 logs\\wewriter.log"
                if on_error:
                    self.events.put(("callback", lambda error=exc: on_error(error)))
                self.events.put(("done", "操作未完成：" + detail))
        threading.Thread(target=work, daemon=True, name="worker").start()

    def poll(self):
        # The poll loop must survive any single bad event; otherwise the
        # progress bar spins forever and the app believes a task is running.
        try:
            self._drain_events()
        except Exception:
            logger.exception("处理界面事件异常")
            self.busy = False
            try:
                self.progress.stop()
            except Exception:
                pass
        if self.closed:
            return
        try:
            if self.root.winfo_exists():
                self.root.after(150, self.poll)
        except tk.TclError:
            pass

    def _drain_events(self):
        try:
            while True:
                kind, text = self.events.get_nowait()
                if kind == "callback":
                    text()
                elif kind == "show":
                    self.root.deiconify()
                    self.root.lift()
                elif kind == "quit":
                    self.close(force=True)
                    if self.closed:
                        return
                elif kind == "done":
                    self.busy = False
                    self.progress.stop()
                    self.progress.pack_forget()
                    self.status.set(redact(text))
                    if getattr(self, 'quit_pending', False):
                        self.close(force=True)
                        if self.closed:
                            return
                        continue
                    self.next_queued()
                    if self.page in ("工作台", "文章库", "公众号"):
                        self.refresh_page()
                else:
                    self.status.set(redact(text))
        except queue.Empty:
            pass

    def schedule_tick(self):
        if not self.demo:
            try:
                config = self.db.config()
                key = due_key(config)
                if key and not self.db.rows("SELECT id FROM jobs WHERE run_key=?", (key,)):
                    job = self.db.job(key, config)
                    if self.busy:
                        # Scheduled and manual tasks share one execution slot.
                        self.db.update_job(job['id'], stage='排队中', message='每日定时任务：排队等待当前任务完成')
                        self.session_queue.append(job['id'])
                    else:
                        self.launch(job, "执行每日定时任务…")
            except Exception:
                logger.exception("定时检查异常")
                self.status.set("定时配置异常，请检查设置")
        self.root.after(30000, self.schedule_tick)

    def minimize(self):
        if not self.tray:
            try:
                import pystray
                from PIL import Image
                icon = Image.open(asset_path("wewriter-cat.png")).convert("RGBA").resize(
                    (64, 64), Image.Resampling.LANCZOS)
                self.tray = pystray.Icon("WeWriter", icon, "公众号工作台", menu=pystray.Menu(
                    pystray.MenuItem("打开工作台", lambda: self.events.put(("show", "")), default=True),
                    pystray.MenuItem("退出软件", lambda: self.events.put(("quit", "")))))
                self.tray.run_detached()
            except Exception:
                self.tray = None
                self.root.iconify()
                self.status.set("托盘不可用，已最小化至任务栏")
                return
        self.cancel_page_timers()
        self.root.withdraw()

    def close(self, force=False):
        if self.editor_dirty or self.settings_dirty:
            if not messagebox.askyesno("未保存修改", "继续将丢弃尚未保存的修改，是否继续？", parent=self.root):
                return
            self.editor_dirty = self.settings_dirty = False
        if not force:
            answer = messagebox.askyesnocancel("关闭窗口", "是否继续在托盘运行？\n\n是：隐藏窗口，继续定时任务。\n否：退出软件，停止定时。", parent=self.root)
            if answer is None:
                return
            if answer:
                self.minimize()
                return
        self.cancel_page_timers()
        if self.busy:
            self.queue_paused = True
            self.engine.stop()
            self.quit_pending = True
            self.status.set("正在停止任务，当前请求结束后软件会自动退出（最长约 100 秒）")
            return
        if self.tray:
            self.tray.stop()
        self.closed = True
        self.root.destroy()
